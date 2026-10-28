"""The formula sheet: how every quantitative BRSR figure is calculated.

Generated from the code that does the calculation (app/brsr/sheet_sources.py
and the derivations in app/brsr/annual_data.py), so what the engagement
manager signs off is always exactly what the platform computes.

Tabs:
  Read me        what this is and how to review it
  Factors        conversion factors, with a sign-off column
  Deviations     where the rule differs from the engagement's Report Mapping
  Figures        every figure: sheet, columns, rule, filters, conversion, where used
  Typed in BRSR  figures with no data-sheet column (typed in the BRSR table)
  Questionnaire  the questions answered in the questionnaire, by department
"""
from __future__ import annotations

import io
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.brsr import formula as F
from app.brsr.annual_data import DERIVE, METRICS, ROW_DERIVE
from app.brsr.sheet_sources import COMPUTED, COMPUTED_NOTES, FACTORS, SOURCED, SOURCES
from app.datasheets.catalog import SHEETS

HOW = {"sum": "Sum of every month in the year", "latest": "Latest month of the year, per location",
       "first": "First month of the year, per location", "mean": "Average of the monthly values"}
ACROSS = {"sum": "locations added", "mean": "average across locations", "max": "largest location figure"}
GREEN = PatternFill("solid", fgColor="26890D")


def _usage() -> dict[str, list[str]]:
    used: dict[str, set[str]] = defaultdict(set)
    from app.brsr.question_bank import QUESTIONS
    for q in QUESTIONS:
        for expr in (q.get("derive") or {}).values():
            for n in F.names_in(expr):
                used[n].add(q["id"])
        for expr in ((q.get("derive_rows") or {}).get("single") or {}).values():
            for n in F.names_in(expr):
                used[n].add(q["id"])
    # computed figures pass their use on to their inputs
    for key, (deps, _fn) in COMPUTED.items():
        for d in deps:
            used[d] |= used.get(key, set())
    return {k: sorted(v) for k, v in used.items()}


def _label(key: str, src) -> str:
    """The annual-form label where there is one; otherwise the first column, readable."""
    if key in METRICS:
        return METRICS[key]["label"]
    p = src.parts[0]
    col = p.cols[0] if isinstance(p.cols[0], str) else p.cols[0][0]
    text = col.replace("_", " ")
    return f"{text} ({p.cat})" if p.cat else text


def _sheet(wb, title, header, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = GREEN
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    return ws


def build() -> bytes:
    used = _usage()
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    for line in (
        ["BRSR formula sheet — how each quantitative figure is calculated"],
        [],
        ["Every quantitative BRSR table is built from the approved data sheets. This workbook is generated from "
         "the code that does the calculation, so it always matches the platform."],
        ["Please review the Factors and Deviations tabs and record agreement in the sign-off columns."],
        [],
        ["Rules that apply everywhere"],
        ["• A row belongs to the financial year (1 April – 31 March) by its own Start_Date and End_Date."],
        ["• Only approved uploads count. If a dataset has several approved versions, the latest approved one is used."],
        ["• If two approved uploads cover the same month, location and category, the more recent approval is used "
         "and a warning is shown."],
        ["• Blank cells in an uploaded sheet count as 0. A figure is blank only when its sheet has no approved data "
         "for the year; the BRSR shows a warning instead of a silent 0."],
        ["• Totals, percentages and intensities are calculated from these figures in the BRSR tables."],
        ["• Previous-year columns come from the previous year's approved sheets; until those exist they are typed."],
    ):
        ws.append(line)
    ws["A1"].font = Font(bold=True, size=13)
    ws["A6"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 140

    _sheet(wb, "Factors", ["Factor", "Value", "Basis", "Needs sign-off", "Agreed (Y/N)", "Comment"],
           [[f["label"], f["value"], f["basis"], "Yes" if f["signoff"] else "No (exact)", "", ""]
            for f in FACTORS.values()], [52, 12, 48, 14, 12, 40])

    dev_rows = []
    for key, src in SOURCES.items():
        if src.deviation:
            p = src.parts[0]
            dev_rows.append([key, METRICS.get(key, {}).get("label", ""), SHEETS[p.sheet]["title"],
                             ", ".join(c if isinstance(c, str) else c[0] for c in p.cols), src.deviation,
                             ", ".join(used.get(key, [])), "", ""])
    dev_rows += [
        ["scope1, scope2, scope3", "Scope 1 / 2 / 3 emissions", "—", "—",
         "Typed in the BRSR (P6.E7, P6.L2) until emission factors are agreed; the sheets give activity data "
         "(fuel, electricity, refrigerant, travel) but no factors.", "P6.E7, P6.L2", "", ""],
        ["w_consumption", "Water consumption", "Water", "withdrawal − discharge",
         "Report Mapping has no consumption column; consumption = total withdrawal − total discharge.",
         "P6.E3, P6.L1b", "", ""],
        ["A.19b", "Exports as % of turnover", "Finance", "Contribution_of_exports_INR_million ÷ revenue",
         "Revenue from operations is used as turnover (Companies Act s.2(91)).", "A.19b", "", ""],
    ]
    _sheet(wb, "Deviations", ["Figure", "Label", "Sheet", "Columns", "What differs and why", "BRSR questions",
                              "Agreed (Y/N)", "Comment"], dev_rows, [24, 30, 24, 40, 70, 22, 12, 40])

    fig_rows = []
    for key, src in SOURCES.items():
        for i, p in enumerate(src.parts):
            cols = []
            for c in p.cols:
                cols.append(c if isinstance(c, str) else f"{c[0]} × {c[1]:g}")
            rule = HOW[p.how] + (f"; {ACROSS[p.across]}" if p.how in ("latest", "first") else "")
            filters = "; ".join(([f"Category = {p.cat}"] if p.cat else []) + [f"{k} = {v}" for k, v in p.where])
            fig_rows.append([key if i == 0 else "", _label(key, src) if i == 0 else "",
                             SHEETS[p.sheet]["title"], " + ".join(cols), rule, filters,
                             f"× {p.scale:g}" if p.scale != 1 else "", src.note if i == 0 else "",
                             ", ".join(used.get(key, [])) if i == 0 else ""])
    for key, (deps, _fn) in COMPUTED.items():
        note = COMPUTED_NOTES.get(key) or COMPUTED_NOTES.get(key.replace("stress_", "")) or \
            COMPUTED_NOTES.get(key[:-1] + "m" if key.startswith("left_") else key, "")
        fig_rows.append([key, METRICS.get(key, {}).get("label", ""), "(calculated)", ", ".join(deps), note, "", "",
                         "", ", ".join(used.get(key, []))])
    _sheet(wb, "Figures", ["Figure", "Label", "Data sheet", "Columns", "Rule", "Filter", "Unit conversion",
                           "Note", "BRSR questions"], fig_rows, [26, 34, 26, 60, 44, 28, 14, 40, 28])

    typed = []
    for qid, spec in DERIVE.items():
        for cell, expr in spec.items():
            names = F.names_in(expr)
            if not names <= SOURCED:
                missing = sorted(names - SOURCED)
                typed.append([qid, cell, ", ".join(METRICS.get(n, {}).get("label", n) for n in missing)])
    _sheet(wb, "Typed in BRSR", ["BRSR question", "Cell", "Figure(s) with no data-sheet column"], typed,
           [16, 30, 70])

    from app.brsr.question_bank import QUESTIONS
    added = {"A.19c", "P2.E2b", "P3.E12", "P5.E11", "P6.L1a"}
    qrows = [[q["id"], q.get("number", ""), (q.get("prompt") or "")[:160], ", ".join(q.get("departments", [])),
              {"qualitative": "Questionnaire", "both": "Questionnaire + data sheets"}[q["source"]],
              "Added (not in the engagement list) — please confirm the department" if q["id"] in added else ""]
             for q in QUESTIONS if q.get("source") in ("qualitative", "both")]
    rows_derived = ", ".join(sorted(ROW_DERIVE))
    qrows.append(["", "", f"Lists built from the sheets: {rows_derived}", "", "", ""])
    _sheet(wb, "Questionnaire", ["Question", "No.", "Question text", "Department(s)", "Answered in", "Note"], qrows,
           [12, 8, 90, 30, 24, 44])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

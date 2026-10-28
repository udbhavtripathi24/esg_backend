"""BRSR Word report generator.

Builds the Business Responsibility and Sustainability Report in SEBI's
layout (Section A, Section B, Principles 1–9 with Essential and
Leadership indicators) from a frozen snapshot of answers.

Every number printed here comes from app/brsr/engine.py — the same
computation the client and reviewer saw on screen — so the document
cannot disagree with the approved data. Formatting rules fixed here
instead of left to whoever types the report:
- percentages to 2 decimals (the reference filing printed
  74.721734036321);
- counts and rupee amounts with Indian digit grouping (35,97,43,580);
- intensities to 4 significant figures, never in scientific notation;
- blank numeric cells print as "-", never as a misleading 0;
- period headers are always filled ("FY 2024-25 (Current FY)") —
  never a "FY ____" placeholder;
- internal review comments do not exist in the snapshot, so they cannot
  leak into the document.
"""
from __future__ import annotations

import io
import os
import shutil
import subprocess
import tempfile
from datetime import datetime
from typing import Any, Optional

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.brsr import engine
from app.brsr.question_bank import CONDITIONAL_REQUIRED, PRINCIPLES, QUESTIONS, SECTIONS, financial_year_labels
from app.core.errors import AppError

FONT = "Calibri"
GREY = RGBColor(0x59, 0x59, 0x59)
RED = RGBColor(0xC0, 0x00, 0x00)
HEADER_FILL = "D9E2D0"
GROUP_FILL = "EEF2EA"


# ---------------------------------------------------------------------------
# Number formatting
# ---------------------------------------------------------------------------

def indian_group(n: int) -> str:
    s = str(abs(n))
    if len(s) > 3:
        head, tail = s[:-3], s[-3:]
        parts = []
        while len(head) > 2:
            parts.insert(0, head[-2:])
            head = head[:-2]
        if head:
            parts.insert(0, head)
        s = ",".join(parts) + "," + tail
    return ("-" if n < 0 else "") + s


def fmt_number(v: Any, display: str = "number") -> str:
    if v is None or v == "":
        return "-"
    if isinstance(v, str):
        return v
    if isinstance(v, bool):
        return "Yes" if v else "No"
    f = float(v)
    if display == "percent":
        return f"{f:.2f}"
    if display == "intensity":
        if f == 0:
            return "0"
        digits = 4
        import math
        mag = math.floor(math.log10(abs(f)))
        decimals = max(0, digits - 1 - mag)
        return f"{f:.{decimals}f}"
    cents = int(round(abs(f) * 100))  # round once, then split — 9.999 becomes 10, not 9.100
    whole, frac = divmod(cents, 100)
    sign = "-" if f < 0 and cents else ""
    return sign + indian_group(whole) + (f".{frac:02d}" if frac else "")


def _display_for(col: dict, row: dict) -> str:
    if col["kind"] == "computed":
        return col.get("display", "number")
    if row.get("formula"):
        return row.get("display", "number")
    return {"percent": "percent"}.get(col["kind"], "number")


def fill_labels(text: str, labels: dict[str, str]) -> str:
    return text.replace("{cy}", labels["cy"]).replace("{py}", labels["py"]).replace("{ppy}", labels["ppy"])


# ---------------------------------------------------------------------------
# Low-level docx helpers
# ---------------------------------------------------------------------------

def _shade(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


USABLE_CM = 26.5  # A4 landscape minus 1.6 cm side margins


def _set_widths(table, widths_cm: list[float]) -> None:
    """Fixed column widths. Word ignores cell widths unless the table has a
    fixed layout and matching <w:gridCol> entries, so set all three."""
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(Cm(w).twips)))
    for row in table.rows:
        # Merged cells appear once per span in row.cells; set the underlying tc widths by grid position.
        for idx, tc in enumerate(row._tr.tc_lst):
            span = int(tc.tcPr.gridSpan.val) if tc.tcPr is not None and tc.tcPr.gridSpan is not None else 1
            start = sum(int(t.tcPr.gridSpan.val) if t.tcPr is not None and t.tcPr.gridSpan is not None else 1
                        for t in row._tr.tc_lst[:idx])
            w = sum(widths_cm[start:start + span])
            tcPr = tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.insert(0, tcW)
            tcW.set(qn("w:w"), str(int(Cm(w).twips)))
            tcW.set(qn("w:type"), "dxa")


def _label_then_equal(n_value_cols: int, label_cm: float) -> list[float]:
    rest = max(USABLE_CM - label_cm, 1.5 * n_value_cols)
    return [label_cm] + [rest / n_value_cols] * n_value_cols


def _repeat_header(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _cell_text(cell, text: str, *, bold: bool = False, size: float = 8, align=None, italic: bool = False,
               color: Optional[RGBColor] = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    if align is not None:
        p.alignment = align
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        run = p.add_run(line)
        run.font.size = Pt(size)
        run.font.name = FONT
        run.bold = bold
        run.italic = italic
        if color is not None:
            run.font.color.rgb = color
        if i < len(lines) - 1:
            run.add_break()


def _para(doc, text: str, *, bold: bool = False, size: float = 10, italic: bool = False, color=None,
          space_after: float = 4, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    r.font.name = FONT
    if color is not None:
        r.font.color.rgb = color
    return p


def _page_number_field(paragraph) -> None:
    run = paragraph.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), tag)
            run._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            run._r.append(it)
    run.font.size = Pt(8)


# ---------------------------------------------------------------------------
# Answer renderers
# ---------------------------------------------------------------------------

class _Renderer:
    def __init__(self, doc, answers: dict[str, dict], labels: dict[str, str], draft: bool):
        self.doc = doc
        self.answers = answers
        self.labels = labels
        self.draft = draft
        self.comp = engine.Computer(answers)
        self.required = engine.required_ids(answers)
        self.conditional = {qid for (qid, _dep, _when) in CONDITIONAL_REQUIRED}

    def question(self, q: dict) -> None:
        number = q["number"] if q["number"] not in ("—", "") else ""
        prompt = fill_labels(q["prompt"], self.labels)
        if q["type"] == "fields" and q.get("numbered"):
            pass  # A.I prints as the numbered table SEBI uses, without a separate prompt line
        else:
            _para(self.doc, f"{number}. {prompt}" if number else prompt, bold=True, size=9.5, space_after=3)

        # Derived view: sheet-built cells and lists (e.g. CSR projects) included
        ans = self.comp.answers.get(q["id"])
        if q["id"] in self.conditional and q["id"] not in self.required and not engine.is_answered(q, ans):
            _para(self.doc, "Not applicable.", size=9.5)
            return
        if q["type"] != "fields" or not q.get("numbered"):
            if not engine.is_answered(q, ans) and not (q["type"] == "table" and ans):
                self._unanswered(q)
                return
        if ans and ans.get("na"):
            _para(self.doc, f"Not applicable. {ans.get('na_reason', '')}".strip(), size=9.5)
            self._note(ans)
            return
        getattr(self, f"_{q['type']}")(q, ans or {})
        self._note(ans or {})

    def _unanswered(self, q: dict) -> None:
        if self.draft:
            _para(self.doc, "[Not yet answered]", italic=True, size=9, color=GREY)
        else:
            _para(self.doc, "Not reported." if not q["required"] else "-", size=9.5)

    def _note(self, ans: dict) -> None:
        if ans.get("note"):
            _para(self.doc, f"Note: {ans['note']}", italic=True, size=8.5, color=GREY)

    # -- simple types ---------------------------------------------------------
    def _text(self, q, ans) -> None:
        for chunk in (ans.get("text") or "").split("\n"):
            if chunk.strip():
                _para(self.doc, chunk.strip(), size=9.5)

    def _yesno(self, q, ans) -> None:
        a = ans.get("answer")
        txt = {"NA": "Not applicable"}.get(a, a or "-")
        if ans.get("details"):
            txt = f"{txt}. {ans['details']}"
        _para(self.doc, txt, size=9.5)
        if ans.get("link"):
            _para(self.doc, f"Web-link: {ans['link']}", size=9, color=GREY)

    def _number(self, q, ans) -> None:
        v = ans.get("value")
        disp = "percent" if q["kind"] == "percent" else "number"
        txt = fmt_number(v, disp)
        if q["kind"] == "percent" and txt != "-":
            txt += "%"
        elif q.get("unit") and txt != "-":
            txt += f" {q['unit']}"
        _para(self.doc, txt, size=9.5)

    def _fields(self, q, ans) -> None:
        f = ans.get("fields") or {}
        numbered = q.get("numbered")
        ncols = 3 if numbered else 2
        t = self.doc.add_table(rows=0, cols=ncols)
        t.style = "Table Grid"
        last_num = None
        for fd in q["fields"]:
            v = f.get(fd["key"])
            if fd["kind"] in ("amount", "count", "number"):
                shown = fmt_number(v) if fd["key"] != "year_incorporated" else (str(int(v)) if v else "-")
            else:
                shown = "-" if v in (None, "") else str(v)
            row = t.add_row().cells
            if numbered:
                num = fd.get("number", "")
                _cell_text(row[0], num if num != last_num else "", size=8.5)
                last_num = num
                _cell_text(row[1], fd["label"] + ":", size=8.5)
                _cell_text(row[2], shown, size=8.5)
            else:
                _cell_text(row[0], fd["label"], size=8.5)
                _cell_text(row[1], shown, size=8.5)
        _set_widths(t, [1.2, 11.0, USABLE_CM - 12.2] if numbered else [11.0, USABLE_CM - 11.0])
        self._spacer()

    def _spacer(self) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)

    # -- tables ---------------------------------------------------------------
    def _table(self, q, ans) -> None:
        cols = q["columns"]
        rows = q["rows"]
        grid = self.comp.grid(q["id"])
        ncols = len(cols) + 1
        t = self.doc.add_table(rows=0, cols=ncols)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        # Group headers (merged)
        for h in q.get("headers") or []:
            hr = t.add_row()
            _repeat_header(hr)
            i = 0
            for cell_def in h:
                span = cell_def["span"]
                cell = hr.cells[i]
                if span > 1:
                    cell = cell.merge(hr.cells[i + span - 1])
                _cell_text(cell, fill_labels(cell_def["label"], self.labels), bold=True,
                           align=WD_ALIGN_PARAGRAPH.CENTER)
                _shade(cell, HEADER_FILL)
                i += span
        # Column labels
        lr = t.add_row()
        _repeat_header(lr)
        _cell_text(lr.cells[0], q.get("row_header") or "", bold=True)
        _shade(lr.cells[0], HEADER_FILL)
        for i, c in enumerate(cols):
            label = fill_labels(c["label"], self.labels)
            if c.get("unit"):
                label += f" ({c['unit']})"
            _cell_text(lr.cells[i + 1], label, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            _shade(lr.cells[i + 1], HEADER_FILL)
        # Body
        for r in rows:
            br = t.add_row()
            if r.get("header"):
                cell = br.cells[0].merge(br.cells[-1])
                _cell_text(cell, r["label"], bold=True)
                _shade(cell, GROUP_FILL)
                continue
            is_total = bool(r.get("formula"))
            _cell_text(br.cells[0], r["label"], bold=is_total)
            vals = grid.get(r["key"], {})
            for i, c in enumerate(cols):
                v = vals.get(c["key"])
                if c["kind"] in engine.NUMERIC_KINDS or c["kind"] == "computed":
                    txt = fmt_number(v, _display_for(c, r))
                    align = WD_ALIGN_PARAGRAPH.RIGHT
                else:
                    txt = "-" if v in (None, "") else str(v)
                    align = None
                _cell_text(br.cells[i + 1], txt, bold=is_total, align=align)
        # Label column: wide for long row labels; text-heavy columns share the rest.
        longest = max((len(r["label"]) for r in rows), default=10)
        label_cm = 4.0 if longest < 30 else (7.0 if longest < 90 else 10.0)
        weights = [3.0 if c["kind"] == "longtext" else 1.6 if c["kind"] == "text" else 1.0 for c in cols]
        rest = USABLE_CM - label_cm
        _set_widths(t, [label_cm] + [rest * w / sum(weights) for w in weights])
        if q.get("unit"):
            _para(self.doc, f"Unit: {q['unit']}", italic=True, size=8, color=GREY, space_after=2)
        self._spacer()

    def _rows(self, q, ans) -> None:
        rws = ans.get("rows") or []
        if not rws:
            _para(self.doc, "Nil" if ans.get("none") else "-", size=9.5)
            return
        cols = q["columns"]
        t = self.doc.add_table(rows=0, cols=len(cols) + 1)
        t.style = "Table Grid"
        hr = t.add_row()
        _repeat_header(hr)
        _cell_text(hr.cells[0], "S. No.", bold=True)
        _shade(hr.cells[0], HEADER_FILL)
        for i, c in enumerate(cols):
            _cell_text(hr.cells[i + 1], c["label"], bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            _shade(hr.cells[i + 1], HEADER_FILL)
        for n, rv in enumerate(rws, start=1):
            br = t.add_row()
            _cell_text(br.cells[0], str(n))
            for i, c in enumerate(cols):
                v = rv.get(c["key"])
                if c["kind"] in engine.NUMERIC_KINDS:
                    txt = fmt_number(v, "percent" if c["kind"] == "percent" else "number")
                    align = WD_ALIGN_PARAGRAPH.RIGHT
                else:
                    txt = "-" if v in (None, "") else str(v)
                    align = None
                _cell_text(br.cells[i + 1], txt, align=align)
        weights = [3.0 if c["kind"] == "longtext" else 1.6 if c["kind"] == "text" else 1.0 for c in cols]
        rest = USABLE_CM - 1.2
        _set_widths(t, [1.2] + [rest * w / sum(weights) for w in weights])
        self._spacer()

    def _principle_grid(self, q, ans) -> None:
        cells = ans.get("cells") or {}
        span = ans.get("span") or {}
        t = self.doc.add_table(rows=0, cols=10)
        t.style = "Table Grid"
        hr = t.add_row()
        _repeat_header(hr)
        _cell_text(hr.cells[0], "Disclosure Questions", bold=True)
        _shade(hr.cells[0], HEADER_FILL)
        for i, p in enumerate(PRINCIPLES):
            _cell_text(hr.cells[i + 1], p, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            _shade(hr.cells[i + 1], HEADER_FILL)
        for g in q["grid_rows"]:
            br = t.add_row()
            _cell_text(br.cells[0], g["label"])
            if g["key"] in span:
                cell = br.cells[1].merge(br.cells[9])
                _cell_text(cell, str(span[g["key"]]))
                continue
            row = cells.get(g["key"]) or {}
            for i, p in enumerate(PRINCIPLES):
                v = row.get(p)
                short = {"Yes": "Y", "No": "N"}.get(v, v) if g["kind"] == "yesno" else v
                _cell_text(br.cells[i + 1], "-" if v in (None, "") else str(short),
                           align=WD_ALIGN_PARAGRAPH.CENTER if g["kind"] == "yesno" else None)
        _set_widths(t, _label_then_equal(9, 8.5))
        self._spacer()


# ---------------------------------------------------------------------------
# Document
# ---------------------------------------------------------------------------

def build_brsr_docx(snapshot: dict, *, draft: bool = False, generated_at: Optional[datetime] = None) -> bytes:
    """``snapshot``: {"financial_year", "company_name", "answers"} — the
    frozen payload from app/brsr/service.py:frozen_snapshot()."""
    fy = snapshot["financial_year"]
    labels = financial_year_labels(fy)
    # Fill every table built from the annual data before rendering.
    answers = engine.apply_derivations(snapshot["answers"])
    company = snapshot.get("company_name") or ""

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.6))
    sec.top_margin = Cm(1.6)
    sec.bottom_margin = Cm(1.4)

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(10)

    # Header / footer
    hp = sec.header.paragraphs[0]
    hp.text = ""
    hr = hp.add_run(f"{company} — Business Responsibility and Sustainability Report, {labels['cy']}")
    hr.font.size = Pt(8)
    hr.font.color.rgb = GREY
    if draft:
        dp = sec.header.add_paragraph()
        dr = dp.add_run("DRAFT — generated from unapproved answers. Not for filing.")
        dr.bold = True
        dr.font.size = Pt(9)
        dr.font.color.rgb = RED
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = fp.add_run("Page ")
    fr.font.size = Pt(8)
    _page_number_field(fp)

    # Title block
    _para(doc, "Business Responsibility and Sustainability Report", bold=True, size=18, space_after=2)
    _para(doc, company, bold=True, size=13, space_after=2)
    _para(doc, f"Financial year {labels['cy'][3:]}", size=11, color=GREY, space_after=2)
    stamp = (generated_at or datetime.utcnow()).strftime("%d %B %Y")
    _para(doc, ("Draft preview" if draft else "Generated") + f" on {stamp} from the BRSR responses "
          + ("entered so far." if draft else "approved in Deloitte VISTA."), size=8.5, italic=True, color=GREY,
          space_after=10)

    r = _Renderer(doc, answers, labels, draft)
    for s in SECTIONS:
        if s.get("kind") == "data":
            continue  # annual data feeds the tables; it is not a BRSR section itself
        code = s["code"]
        heading = s["title"] if code in ("A", "B") else f"{s['title']}: {s['statement']}"
        if code == "P1":
            h0 = doc.add_heading("SECTION C: PRINCIPLE WISE PERFORMANCE DISCLOSURE", level=1)
            for run in h0.runs:
                run.font.name = FONT
        h = doc.add_heading(heading.upper() if code in ("A", "B") else heading, level=1 if code in ("A", "B") else 2)
        for run in h.runs:
            run.font.name = FONT
        if s.get("intro"):
            _para(doc, s["intro"], size=9.5, italic=True)
        current_group = None
        for q in QUESTIONS:
            if q["section"] != code or q.get("helper"):
                continue
            if q.get("group") and q["group"] != current_group:
                current_group = q["group"]
                gp = _para(doc, current_group, bold=True, size=10.5, space_after=4)
                gp.paragraph_format.space_before = Pt(6)
            r.question(q)

    from app.brsr.sheet_metrics import SHEET_KEY
    if (answers.get(SHEET_KEY) or {}).get("basis"):
        _data_annex(doc, answers[SHEET_KEY], labels)
    if snapshot.get("assurance"):
        _assurance_annex(doc, snapshot["assurance"], labels)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _data_annex(doc, sheet: dict, labels: dict) -> None:
    """Which approved data-sheet uploads the quantitative tables were built from."""
    from app.datasheets.catalog import SHEETS
    doc.add_page_break()
    h = doc.add_heading("ANNEXURE: DATA SOURCES", level=1)
    for run in h.runs:
        run.font.name = FONT
    _para(doc, "The quantitative disclosures are calculated from the following data sheets, as approved in "
               "Deloitte VISTA. Each figure can be traced to the uploaded row it came from.", size=9.5)
    cols = ["Data sheet", "Upload", "Version", f"Months ({labels['cy']})", f"Months ({labels['py']})"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        _cell_text(t.rows[0].cells[i], c, bold=True)
        _shade(t.rows[0].cells[i], "E8F3E3")
    _repeat_header(t.rows[0])
    cy, py = sheet.get("cy_coverage") or {}, sheet.get("py_coverage") or {}
    for b in sorted(sheet.get("basis") or [], key=lambda x: SHEETS.get(x["sheet"], {}).get("title", x["sheet"])):
        cells = t.add_row().cells
        values = [SHEETS.get(b["sheet"], {}).get("title", b["sheet"]), b.get("dataset_public_id") or "",
                  f"v{b.get('version_number')}", str((cy.get(b["sheet"]) or {}).get("months", 0)),
                  str((py.get(b["sheet"]) or {}).get("months", 0))]
        for i, v in enumerate(values):
            _cell_text(cells[i], v)
    _set_widths(t, [6.0, 5.0, 2.0, 3.0, 3.0])


OPINION_TEXT = {"unmodified": "Unmodified", "qualified": "Qualified", "adverse": "Adverse",
                "disclaimer": "Disclaimer of conclusion"}


def _assurance_annex(doc, a: dict, labels: dict) -> None:
    """Independent assurance and the sign-off chain, from a signed engagement."""
    doc.add_page_break()
    h = doc.add_heading("ANNEXURE: INDEPENDENT ASSURANCE AND SIGN-OFF", level=1)
    for run in h.runs:
        run.font.name = FONT
    signed = (a.get("signed_at") or "")[:10]
    _para(doc, f"{a['level_label']} was obtained from {a['provider_name']} under {a['standard']}. "
               f"Conclusion: {OPINION_TEXT.get(a['opinion'], a['opinion'])}. Signed by {a.get('signed_by') or ''} "
               f"on {signed}. The full assurance statement is issued separately by the assurance provider.",
          size=9.5)
    if a.get("opinion_basis"):
        _para(doc, f"Basis for conclusion: {a['opinion_basis']}", size=9.5, italic=True)
    _para(doc, "Sign-off register — who prepared, reviewed, approved and assured each section", bold=True,
          size=10.5, space_after=4)
    cols = ["Section", "Prepared by", "Submitted by", "Internal review", "Approved by", "Tested by",
            "Workpaper reviewed by"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        _cell_text(t.rows[0].cells[i], c, bold=True)
        _shade(t.rows[0].cells[i], "E8F3E3")
    _repeat_header(t.rows[0])

    def who(name, at):
        return f"{name}\n{(at or '')[:10]}" if name else "–"

    for r in a.get("register") or []:
        cells = t.add_row().cells
        prepared = list(r.get("prepared_by") or []) + [f"(platform: {x})" for x in r.get("system_sources") or []]
        values = [r["section"], "; ".join(prepared) or "–",
                  who(r.get("submitted_by"), r.get("submitted_at")),
                  who(r.get("internal_reviewed_by"), r.get("internal_reviewed_at")),
                  who(r.get("approved_by"), r.get("approved_at")), who(r.get("tested_by"), r.get("tested_at")),
                  who(r.get("workpaper_reviewed_by"), r.get("workpaper_reviewed_at"))]
        for i, v in enumerate(values):
            _cell_text(cells[i], v)
    _set_widths(t, [2.0, 5.0, 3.6, 3.6, 3.6, 3.6, 3.6])
    _para(doc, f"Fingerprint of the assured BRSR answers (SHA-256): {a.get('content_hash')}", size=7.5,
          color=GREY)


def validate_brsr_docx(data: bytes) -> list[str]:
    """Technical validation of a generated file before it is stored."""
    errors = []
    try:
        d = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        return [f"document does not open: {exc}"]
    text = "\n".join(p.text for p in d.paragraphs)
    for needle in ("SECTION A", "SECTION B", "SECTION C", "Principle 9"):
        if needle.lower() not in text.lower():
            errors.append(f"missing heading '{needle}'")
    if "FY ____" in text:
        errors.append("unfilled financial-year placeholder")
    if len(d.tables) < 40:
        errors.append(f"only {len(d.tables)} tables rendered")
    return errors


def pdf_available() -> bool:
    return shutil.which("soffice") is not None or shutil.which("libreoffice") is not None


def docx_to_pdf(data: bytes, timeout: int = 120) -> bytes:
    """Convert with headless LibreOffice. Raises a clear 503 if the server
    image does not include it, instead of failing obscurely."""
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        raise AppError("pdf_unavailable",
                       "PDF export needs LibreOffice in the server image. Download the Word file instead.", 503)
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, "report.docx")
        with open(src, "wb") as fh:
            fh.write(data)
        env = dict(os.environ, HOME=tmp)  # LibreOffice needs a writable profile dir
        proc = subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", tmp, src],
                              capture_output=True, timeout=timeout, env=env)
        out = os.path.join(tmp, "report.pdf")
        if proc.returncode != 0 or not os.path.exists(out):
            raise AppError("pdf_failed", "PDF conversion failed. Download the Word file instead.", 500)
        with open(out, "rb") as fh:
            return fh.read()

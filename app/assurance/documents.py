"""Assurance documents: the statement, the sign-off register and the pack.

At signing, the statement (Word) and the register (Excel) are generated
once and written to storage with their SHA-256; downloads after that serve
those stored files, never a regenerated copy, so the document a reader
gets is the document that was signed.

The assurance pack is a zip assembled on request: the two signed documents,
every evidence file arranged by section, and a MANIFEST listing each file's
SHA-256. Evidence is re-hashed while packing and compared with the checksum
recorded at upload, so tampering in storage would show in the manifest.
"""
from __future__ import annotations

import hashlib
import io
import re
import tempfile
import zipfile
from datetime import datetime
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlmodel import Session, select

from app.models.assurance import AssuranceEngagement, AssuranceFinding, AssuranceRequest
from app.models.company import Company
from app.models.evidence import EvidenceFile
from app.models.user import User
from app.storage import get_storage

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
BRSR_CRITERIA = ("the Business Responsibility and Sustainability Reporting format prescribed by the Securities and "
                 "Exchange Board of India (circular SEBI/HO/CFD/CFD-SEC-2/P/CIR/2023/122 and related guidance)")
OPINION_LABEL = {"unmodified": "Unmodified", "qualified": "Qualified", "adverse": "Adverse",
                 "disclaimer": "Disclaimer of conclusion"}


def _fmt_dt(v) -> str:
    if not v:
        return ""
    if isinstance(v, str):
        return v
    return v.strftime("%d %b %Y %H:%M UTC") if isinstance(v, datetime) else v.strftime("%d %b %Y")


def _company(session: Session, eng: AssuranceEngagement) -> str:
    c = session.get(Company, eng.company_id)
    return c.name if c else "the Company"


# ---------------------------------------------------------------------------
# Statement
# ---------------------------------------------------------------------------

def conclusion_paragraph(eng: AssuranceEngagement, company: str) -> str:
    subject = f"the Business Responsibility and Sustainability Report of {company} for FY {eng.financial_year}"
    if eng.opinion == "disclaimer":
        return (f"Because of the significance of the matter(s) described in the basis for disclaimer, we were unable "
                f"to obtain sufficient appropriate evidence to form a conclusion. Accordingly, we do not express a "
                f"conclusion on {subject}.")
    if eng.opinion == "adverse":
        return (f"Because of the significance of the matter(s) described in the basis for adverse conclusion, "
                f"{subject} is not prepared, in all material respects, in accordance with the criteria.")
    limited = eng.level == "limited"
    except_for = ("Except for the effects of the matter(s) described in the basis for qualified conclusion, "
                  if eng.opinion == "qualified" else "")
    if limited:
        lead = (except_for + "based on" if except_for else "Based on")
        return (f"{lead} the procedures we have performed and the evidence we have obtained, nothing has come to our "
                f"attention that causes us to believe that {subject} is not prepared, in all material respects, in "
                f"accordance with the criteria.")
    lead = except_for + ("in" if except_for else "In")
    return f"{lead} our opinion, {subject} is prepared, in all material respects, in accordance with the criteria."


def default_statement(session: Session, eng: AssuranceEngagement, signer: User) -> str:
    """Plain-text statement used when the partner does not supply their own."""
    company = _company(session, eng)
    n_req = len(session.exec(select(AssuranceRequest.id).where(AssuranceRequest.engagement_id == eng.id)).all())
    n_find = len(session.exec(select(AssuranceFinding.id).where(AssuranceFinding.engagement_id == eng.id)).all())
    level = "limited" if eng.level == "limited" else "reasonable"
    parts = [
        f"To the Board of Directors of {company}",
        f"We have undertaken a {level} assurance engagement on the Business Responsibility and Sustainability Report "
        f"(BRSR) of {company} for the financial year {eng.financial_year}, covering the following sections: "
        f"{', '.join(eng.scope)}.",
        f"Criteria: {BRSR_CRITERIA}.",
        f"Management's responsibility: management is responsible for preparing the BRSR in accordance with the "
        f"criteria, and for the internal controls, records and evidence supporting it.",
        f"Our responsibility: to express a {level} assurance conclusion based on the procedures we performed, in "
        f"accordance with {eng.standard}. We complied with the independence and other ethical requirements "
        f"applicable to this engagement; every member of the engagement team declared independence of the "
        f"information assured.",
        f"Procedures performed included enquiries of the persons responsible for the data, analytical procedures, "
        f"and examination of supporting evidence on a test basis ({n_req} information request(s) raised, "
        f"{n_find} finding(s) recorded and resolved).",
    ]
    if eng.level == "limited":
        parts.append("The procedures performed in a limited assurance engagement vary in nature and timing from, and "
                     "are less in extent than, a reasonable assurance engagement; consequently the level of assurance "
                     "obtained is substantially lower.")
    if eng.opinion != "unmodified" and eng.opinion_basis:
        parts.append(f"Basis for {OPINION_LABEL[eng.opinion].lower()} conclusion: {eng.opinion_basis}")
    parts.append(f"Conclusion: {conclusion_paragraph(eng, company)}")
    parts.append(f"For {eng.provider_name}\n{signer.name}\nEngagement Partner\n"
                 f"{(eng.signed_at or datetime.utcnow()).strftime('%d %B %Y')}")
    return "\n\n".join(parts)


def build_statement_docx(session: Session, eng: AssuranceEngagement) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    h = doc.add_heading("Independent Assurance Statement", level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"{eng.level.capitalize()} assurance · {OPINION_LABEL.get(eng.opinion or '', '')} conclusion · "
              f"FY {eng.financial_year}").italic = True
    for para in (eng.statement_text or "").split("\n\n"):
        doc.add_paragraph(para)
    doc.add_paragraph()
    t = doc.add_table(rows=0, cols=2)
    t.style = "Light Grid Accent 1"
    for k, v in (("Assurance provider", eng.provider_name), ("Standard", eng.standard),
                 ("Level", f"{eng.level.capitalize()} assurance"),
                 ("Conclusion", OPINION_LABEL.get(eng.opinion or "", "")),
                 ("Signed at", _fmt_dt(eng.signed_at)),
                 ("Fingerprint of the assured BRSR (SHA-256)", eng.signed_content_hash or "")):
        r = t.add_row().cells
        r[0].text, r[1].text = k, v
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Register (Excel)
# ---------------------------------------------------------------------------

HEADER_FILL = PatternFill("solid", fgColor="26890D")


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="top", wrap_text=True)
    for r in rows:
        ws.append([_fmt_dt(v) if isinstance(v, datetime) else v for v in r])
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1] or "")) for r in rows[:200]] + [8])
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 60)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    return ws


def build_register_xlsx(session: Session, eng: AssuranceEngagement) -> bytes:
    from app.assurance import service

    d_names = service.brsr_names
    company = _company(session, eng)
    team = service.active_members(session, eng)
    names = d_names(session, {m.user_id for m in team} | {eng.signed_by, eng.eqr_cleared_by, eng.created_by})
    wb = Workbook()
    status = eng.status if eng.status != "signed" else "signed"
    _sheet(wb, "Summary", ["Item", "Value"], [
        ["Company", company], ["Financial year", eng.financial_year], ["Assurance provider", eng.provider_name],
        ["Level", f"{eng.level.capitalize()} assurance"], ["Standard", eng.standard], ["Status", status],
        ["Conclusion", OPINION_LABEL.get(eng.opinion or "", "") if eng.opinion else ""],
        ["Basis for conclusion", eng.opinion_basis or ""],
        ["Engagement quality review cleared by", names.get(eng.eqr_cleared_by, "")],
        ["Engagement quality review cleared at", eng.eqr_cleared_at],
        ["Signed by (engagement partner)", names.get(eng.signed_by, "")], ["Signed at", eng.signed_at],
        ["Fingerprint of the assured BRSR (SHA-256)", eng.signed_content_hash or "(not yet signed)"],
        ["Register generated at", datetime.utcnow()],
    ], first=True)
    _sheet(wb, "Sign-off register", [
        "Section", "Title", "Prepared by", "Last edited", "Evidence files", "Submitted by", "Submitted at",
        "Internal review by", "Internal review at", "Approved by (Deloitte)", "Approved at", "Tested by (assurance)",
        "Tested at", "Workpaper reviewed by", "Workpaper reviewed at", "Conclusion"],
        [[r["section"], r["title"], "; ".join(r["prepared_by"] + [f"(platform: {x})" for x in r.get("system_sources", [])]),
          r["last_edited_at"], r["evidence_files"],
          r["submitted_by"], r["submitted_at"], r["internal_reviewed_by"], r["internal_reviewed_at"],
          r["approved_by"], r["approved_at"], r["tested_by"], r["tested_at"], r["workpaper_reviewed_by"],
          r["workpaper_reviewed_at"], (r["conclusion"] or "").replace("_", " ")]
         for r in (eng.signed_register or service.register_rows(session, eng))])
    _sheet(wb, "Team & independence", ["Name", "Role", "Added at", "Independence declared at"],
           [[names.get(m.user_id), m.role, m.added_at, m.independence_declared_at] for m in team])
    reqs = session.exec(select(AssuranceRequest).where(AssuranceRequest.engagement_id == eng.id)
                        .order_by(AssuranceRequest.requested_at)).all()
    findings = session.exec(select(AssuranceFinding).where(AssuranceFinding.engagement_id == eng.id)
                            .order_by(AssuranceFinding.raised_at)).all()
    evid = session.exec(select(EvidenceFile).where(EvidenceFile.response_set_id == eng.response_set_id,
                                                   EvidenceFile.deleted_at.is_(None))
                        .order_by(EvidenceFile.section_code, EvidenceFile.uploaded_at)).all()
    more = d_names(session, {x for r in reqs for x in (r.requested_by, r.responded_by, r.closed_by)}
                   | {x for f in findings for x in (f.raised_by, f.responded_by, f.resolved_by)}
                   | {e.uploaded_by for e in evid})
    _sheet(wb, "Information requests", ["Section", "Question", "Request", "Requested by", "Requested at", "Due",
                                        "Response", "Responded by", "Responded at", "Status", "Closed by",
                                        "Closed at", "Times sent back"],
           [[r.section_code, r.question_id, r.title, more.get(r.requested_by), r.requested_at,
             r.due_date.isoformat() if r.due_date else "", r.response_text, more.get(r.responded_by),
             r.responded_at, r.status, more.get(r.closed_by), r.closed_at, len(r.history)] for r in reqs])
    _sheet(wb, "Findings", ["Section", "Question", "Severity", "Finding", "Raised by", "Raised at",
                            "Management response", "Responded by", "Outcome", "Resolved by", "Resolved at",
                            "Resolution note"],
           [[f.section_code, f.question_id, f.severity.replace("_", " "), f.description, more.get(f.raised_by),
             f.raised_at, f.management_response, more.get(f.responded_by), f.status, more.get(f.resolved_by),
             f.resolved_at, f.resolution_note] for f in findings])
    _sheet(wb, "Evidence", ["Section", "Question", "File", "Uploaded by", "Uploaded at", "For assurance request",
                            "SHA-256"],
           [[e.section_code, e.question_id, e.original_filename, more.get(e.uploaded_by), e.uploaded_at,
             "Yes" if e.assurance_request_id else "", e.sha256_checksum] for e in evid])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Signing and retrieval
# ---------------------------------------------------------------------------

def _base_key(eng: AssuranceEngagement) -> str:
    return f"companies/{eng.company_id}/assurance/{eng.public_id}"


def _filename(eng: AssuranceEngagement, kind: str, ext: str) -> str:
    return f"BRSR_FY{eng.financial_year}_{kind}.{ext}"


def store_signed_documents(session: Session, eng: AssuranceEngagement) -> None:
    storage = get_storage()
    statement = build_statement_docx(session, eng)
    stored = storage.put(f"{_base_key(eng)}/{_filename(eng, 'assurance_statement', 'docx')}",
                         io.BytesIO(statement), mime_type=DOCX_MIME)
    eng.statement_storage_key, eng.statement_sha256 = stored.key, stored.sha256_checksum
    register = build_register_xlsx(session, eng)
    stored = storage.put(f"{_base_key(eng)}/{_filename(eng, 'signoff_register', 'xlsx')}",
                         io.BytesIO(register), mime_type=XLSX_MIME)
    eng.register_storage_key, eng.register_sha256 = stored.key, stored.sha256_checksum


def artifact_for_key(session: Session, key: str) -> Optional[tuple[str, str]]:
    """(mime, filename) when ``key`` is a signed assurance document."""
    eng = session.exec(select(AssuranceEngagement).where(
        (AssuranceEngagement.statement_storage_key == key) | (AssuranceEngagement.register_storage_key == key))).first()
    if not eng:
        return None
    if key == eng.statement_storage_key:
        return DOCX_MIME, _filename(eng, "assurance_statement", "docx")
    return XLSX_MIME, _filename(eng, "signoff_register", "xlsx")


def signed_document_url(eng: AssuranceEngagement, kind: str) -> dict:
    key = eng.statement_storage_key if kind == "statement" else eng.register_storage_key
    return {"url": get_storage().signed_url(key, expires_in=300), "expires_in": 300,
            "filename": _filename(eng, "assurance_statement" if kind == "statement" else "signoff_register",
                                  "docx" if kind == "statement" else "xlsx"),
            "sha256": eng.statement_sha256 if kind == "statement" else eng.register_sha256}


def _read(key: str) -> bytes:
    stream = get_storage().get(key)
    try:
        return stream.read()
    finally:
        close = getattr(stream, "close", None)
        if close:
            close()


def build_pack(session: Session, eng: AssuranceEngagement):
    """A file object (positioned at 0) holding the assurance pack zip."""
    from app.assurance import service

    out = tempfile.SpooledTemporaryFile(max_size=64 * 1024 * 1024)
    manifest = [f"Assurance pack — {_company(session, eng)} — BRSR FY {eng.financial_year}",
                f"Provider: {eng.provider_name} · {eng.level.capitalize()} assurance · "
                f"{OPINION_LABEL.get(eng.opinion or '', '')} conclusion",
                f"Signed at: {_fmt_dt(eng.signed_at)}",
                f"Fingerprint of the assured BRSR (SHA-256): {eng.signed_content_hash}",
                f"Re-assurance required (BRSR changed since signing): "
                f"{'YES' if service.is_stale(session, eng) else 'no'}",
                "", "SHA-256  path  (evidence: checked against the checksum recorded at upload)"]
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as z:
        for key, sha, name in ((eng.statement_storage_key, eng.statement_sha256,
                                _filename(eng, "assurance_statement", "docx")),
                               (eng.register_storage_key, eng.register_sha256,
                                _filename(eng, "signoff_register", "xlsx"))):
            data = _read(key)
            actual = hashlib.sha256(data).hexdigest()
            z.writestr(name, data)
            manifest.append(f"{actual}  {name}" + ("" if actual == sha else "  ** DIFFERS FROM SIGNED COPY **"))
        used: set[str] = set()
        for ev in session.exec(select(EvidenceFile).where(EvidenceFile.response_set_id == eng.response_set_id,
                                                          EvidenceFile.deleted_at.is_(None))
                               .order_by(EvidenceFile.section_code, EvidenceFile.uploaded_at)).all():
            data = _read(ev.storage_key)
            actual = hashlib.sha256(data).hexdigest()
            safe = re.sub(r"[^\w.\- ]", "_", ev.original_filename)
            path = f"evidence/{ev.section_code}/{ev.public_id[-6:]}_{safe}"
            while path in used:
                path = path + "_"
            used.add(path)
            z.writestr(path, data)
            manifest.append(f"{actual}  {path}" + ("" if actual == ev.sha256_checksum
                                                    else "  ** DOES NOT MATCH THE CHECKSUM RECORDED AT UPLOAD **"))
        z.writestr("MANIFEST.txt", "\n".join(manifest) + "\n")
    out.seek(0)
    return out

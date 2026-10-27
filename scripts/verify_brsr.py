"""BRSR end-to-end verification against the real API (the endpoints the UI calls).

Covers the short Annual data form (bank v2): figures entered once per
department area, the BRSR tables built from them, validation at the source,
review, reporting gates and the draft preview.

Usage (from the esg-backend repo root, with the venv active):
    python scripts/verify_brsr.py http://127.0.0.1:8001

Prints PASS/FAIL per check and exits non-zero if anything fails.
Expects the throwaway users created by scripts/verify_brsr_setup.py.
"""
import io
import sys

import httpx
from docx import Document

ROOT = (sys.argv[1] if len(sys.argv) > 1 else "http://127.0.0.1:8001").rstrip("/")
BASE = ROOT + "/api/v1"
PW = "verify-pass-123"
results = []


def check(name, ok, detail=""):
    results.append(bool(ok))
    print(f"[{'PASS' if ok else 'FAIL'}] {name}" + (f"  ->  {detail}" if detail else ""))


def login(email):
    r = httpx.post(f"{BASE}/auth/login", data={"username": email, "password": PW})
    r.raise_for_status()
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


client = login("verify.clientadmin@brsr.test")
reviewer = login("verify.deloitte@brsr.test")

# ------------------------------------------------------------------ 0. health
print("0. Readiness")
r = httpx.get(f"{ROOT}/ready")
check("/ready reports database ok and migrations up to date",
      r.status_code == 200 and r.json().get("migrations") == "up_to_date", r.text)

# ---------------------------------------------------------------- setup set
r = httpx.post(f"{BASE}/brsr/response-sets", json={"financial_year": "2024-25"}, headers=client)
if r.status_code == 409:
    sid = httpx.get(f"{BASE}/brsr/response-sets", headers=client).json()[0]["public_id"]
else:
    r.raise_for_status()
    sid = r.json()["public_id"]
S = f"{BASE}/brsr/response-sets/{sid}"
print(f"response set: {sid}\n")
bank = httpx.get(f"{BASE}/brsr/question-bank", headers=client).json()
ov = httpx.get(S, headers=client).json()
data_areas = [s["code"] for s in ov["sections"] if s.get("kind") == "data"]
check("questionnaire has the Annual data areas first, then the 11 SEBI sections",
      len(data_areas) == 13 and len(ov["sections"]) == 24 and ov["sections"][0]["code"] == "WF",
      f"{len(data_areas)} data areas, {len(ov['sections'])} sections")
check("every question is routed to a department",
      all(q.get("departments") for q in bank["questions"] if not q.get("helper")))


def save(answers, headers=client):
    return httpx.put(f"{S}/answers", json={"answers": answers}, headers=headers, timeout=60)


def answer(code, qid):
    return httpx.get(f"{S}/sections/{code}", headers=client).json()["answers"][qid]


NO_WORKERS = {"wrk_perm_m": 0, "wrk_perm_f": 0, "wrk_oth_m": 0, "wrk_oth_f": 0}

# ---------------------------------------- 1. validation at the source (Annual data)
print("\n1. Annual data: validation and the tables built from it")
r = save({
    "D.WF": {"value": {"values": {"emp_perm_m": 2551, "emp_perm_f": 863, "emp_oth_m": 0, "emp_oth_f": 0,
                                  **NO_WORKERS}}, "revision": answer("WF", "D.WF")["revision"]},
    "D.TR": {"value": {"values": {"rev_emp_m": 2862, "rev_emp_f": 846}},
             "revision": answer("TR", "D.TR")["revision"]},
})
body = r.json()
check("annual figures saved", r.status_code == 200, f"HTTP {r.status_code}")
msgs = [i["message"] for i in body["answers"]["D.TR"]["issues"]]
check("error at the source: reviews cannot exceed headcount",
      any("2,862" in m and "cannot exceed" in m for m in msgs), msgs[0] if msgs else "no issues")
check("Training area not submittable while the error exists", body["sections"]["TR"]["ready_to_submit"] is False)
p3 = answer("P3", "P3.E9")
check("P3 Q9 table is built from the annual data (112.19% shows the mistake)",
      round(p3["computed"]["emp_male"]["cy_pct"], 2) == 112.19, str(p3["computed"]["emp_male"]))

wf = answer("WF", "D.WF")
r = save({"D.WF": {"value": {**wf["value"], "values": {**wf["value"]["values"], "emp_oth_m": 406, "emp_oth_f": 83}},
                   "revision": wf["revision"]}})
check("headcount corrected (contract staff added)", r.status_code == 200, f"HTTP {r.status_code}")
tr = answer("TR", "D.TR")
check("error cleared once the headcount is complete", not any("cannot exceed" in i["message"] for i in tr["issues"]))
p3 = answer("P3", "P3.E9")
check("P3 Q9 now 96.79% — computed, never typed", round(p3["computed"]["emp_male"]["cy_pct"], 2) == 96.79,
      str(round(p3["computed"]["emp_male"]["cy_pct"], 2)))
a20 = answer("A", "A.20a")
check("Section A Q20 uses the same headcount (3,903 employees)",
      a20["computed"]["employees_total"]["total"] == 3903, str(a20["computed"]["employees_total"]))

r = save({"P3.E9": {"value": {"cells": {"emp_male": {"cy_total": 1, "cy_reviewed": 1}}},
                    "revision": p3["revision"]}})
p3b = answer("P3", "P3.E9")
check("built tables cannot be overwritten by typing", r.status_code == 200 and
      round(p3b["computed"]["emp_male"]["cy_pct"], 2) == 96.79, str(p3b["computed"]["emp_male"]))

wf = answer("WF", "D.WF")
check("worker follow-up questions hidden once workers = 0", len(wf["hidden_groups"]) >= 1, str(wf["hidden_groups"]))

r = save({"D.WF": {"value": {"values": {"emp_perm_m": 1}}, "revision": 1}})
check("stale edit refused with 409 (no silent overwrite)", r.status_code == 409, f"HTTP {r.status_code}")

# --------------------------------------------------------------- 2. persistence
print("\n2. Persistence (fresh read, same as a page refresh)")
wf = answer("WF", "D.WF")
check("stored headcount is 2551 / 406", wf["value"]["values"].get("emp_perm_m") == 2551 and
      wf["value"]["values"].get("emp_oth_m") == 406, str(wf["value"]["values"]))
hist = httpx.get(f"{S}/answers/D.WF/history", headers=client).json()
check("change history recorded", len(hist) >= 2, f"{len(hist)} revisions")

# ------------------------------------------------------------------ 3. review
print("\n3. Submit Principle 4 and review as Deloitte")
r = save({
    "P4.E1": {"value": {"text": "Stakeholders are identified through an annual materiality exercise."},
              "revision": answer("P4", "P4.E1")["revision"]},
    "P4.E2": {"value": {"rows": [{"group": "Employees", "vulnerable": "No", "channels": "Email, town halls",
                                  "frequency": "Quarterly", "purpose": "Career growth and well-being"}]},
              "revision": answer("P4", "P4.E2")["revision"]},
})
check("P4 answers saved", r.status_code == 200 and r.json()["sections"]["P4"]["ready_to_submit"], f"HTTP {r.status_code}")
r = httpx.post(f"{S}/sections/P4/submit", headers=client)
check("P4 submitted", r.status_code == 200 and r.json()["status"] == "submitted", f"HTTP {r.status_code}")
r = save({"P4.E1": {"value": {"text": "edit after submit"}, "revision": answer("P4", "P4.E1")["revision"]}})
check("submitted section is locked", r.status_code == 409 and r.json()["error"]["code"] == "section_locked",
      f"HTTP {r.status_code}")
q = httpx.get(f"{BASE}/brsr/review-queue", headers=reviewer).json()
row = next((x for x in q if x["response_set_id"] == sid and x["section"] == "P4"), None)
check("P4 appears in reviewer's queue with company + submitter", row is not None,
      f"{row['company_name']} / submitted by {row['submitted_by']}" if row else "not in queue")
r = httpx.post(f"{S}/sections/P4/decide", json={"decision": "approved", "note": "Verified."}, headers=reviewer)
check("reviewer approves P4", r.status_code == 200 and r.json()["status"] == "approved", f"HTTP {r.status_code}")

# --------------------------------------------------------------- 4. reporting
print("\n4. Reporting readiness and draft preview")
ov = httpx.get(S, headers=client).json()
n = sum(1 for s in ov["sections"] if s["status"] == "approved")
check("readiness shows 1/24 approved", n == 1 and len(ov["sections"]) == 24, f"{n}/{len(ov['sections'])}")
r = httpx.post(f"{BASE}/reports", json={"reporting_period_start": "2024-04-01", "reporting_period_end": "2025-03-31",
                                        "report_type": "brsr"}, headers=client)
check("final report refused until every section is approved",
      r.status_code == 422 and r.json()["error"]["code"] == "brsr_not_ready", f"HTTP {r.status_code}")
r = httpx.get(f"{S}/preview", params={"fmt": "docx"}, headers=client, timeout=120)
ok = r.status_code == 200 and r.content[:2] == b"PK"
check("draft preview downloads a .docx", ok, f"HTTP {r.status_code}, {len(r.content)} bytes")
if ok:
    doc = Document(io.BytesIO(r.content))
    header = " ".join(p.text for p in doc.sections[0].header.paragraphs)
    text = " ".join(p.text for p in doc.paragraphs)
    cells_text = " ".join(c.text for t in doc.tables for row_ in t.rows for c in row_.cells)
    check("preview header is watermarked DRAFT", "DRAFT" in header and "Not for filing" in header)
    check("preview prints the built 96.79 and 3,903", "96.79" in cells_text and "3,903" in cells_text)
    check("Annual data itself is not printed (SEBI format only)", "Workforce & diversity" not in text + cells_text)
r = httpx.get(f"{S}/preview", params={"fmt": "pdf"}, headers=client, timeout=180)
if r.status_code == 503:
    print("[INFO] PDF preview: LibreOffice not installed on this machine (expected locally; the Docker image includes it)")
else:
    check("draft PDF downloads", r.status_code == 200 and r.content[:4] == b"%PDF", f"HTTP {r.status_code}")

print(f"\n{sum(results)}/{len(results)} checks passed")
sys.exit(0 if all(results) else 1)

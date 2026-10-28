"""Evidence, two-tier review and independent assurance of the BRSR.

The journey test walks a real engagement end to end through the API:
team selection with the independence check, fieldwork (workpapers with
four-eyes review, an information request answered with evidence, findings
corrected and left uncorrected), the quality review, the partner's
signature, the signed documents and the assurance pack, the BRSR report
annex, and the "needs re-assurance" flag after the data changes.
"""
import hashlib
import io
import zipfile

import pytest
from docx import Document
from openpyxl import load_workbook
from sqlmodel import select

from app.models.company import Company
from app.models.consultant_assignment import ConsultantAssignment
from tests import brsr_fixture
from tests.conftest_helpers import auth, bootstrap, make_company, make_user
from tests.test_brsr import API, _approve_all, _approve_answers, _create, _fill_all, _save, _submit_all
from tests.test_layer1_kpi import _set_storage

A = "/api/v1/assurance"
PDF = b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF\n"
PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64


def _world(session):
    org = bootstrap(session)
    co = make_company(session, org, "Nuvama Wealth")
    other = make_company(session, org, "Other Co")
    w = {
        "org": org, "co": co,
        "uploader": make_user(session, "up@client.com", "client", "Client Uploader", company=co),
        "uploader2": make_user(session, "up2@client.com", "client", "Client Uploader", company=co),
        "admin": make_user(session, "admin@client.com", "client", "Client Administrator", company=co),
        "approver": make_user(session, "appr@client.com", "client", "Client Approver", company=co),
        "outsider": make_user(session, "x@other.com", "client", "Client Administrator", company=other),
        "reviewer": make_user(session, "rev@deloitte.com", "deloitte", "Consultant", org=org),
        "dadmin": make_user(session, "dadmin@deloitte.com", "deloitte", "Administrator", org=org),
        "partner": make_user(session, "partner@assure.com", "deloitte", "Assurance Partner", org=org),
        "eqr": make_user(session, "eqr@assure.com", "deloitte", "Assurance Partner", org=org),
        "prac": make_user(session, "prac@assure.com", "deloitte", "Assurance Practitioner", org=org),
        "prac2": make_user(session, "prac2@assure.com", "deloitte", "Assurance Practitioner", org=org),
    }
    session.add(ConsultantAssignment(company_id=co.id, consultant_user_id=w["reviewer"].id, is_active=True))
    session.commit()
    brsr_fixture.seed_sheet_data(session, co, w["uploader"])
    return w


def _upload(client, user, set_id, code, name, data, question_id=None):
    files = {"file": (name, io.BytesIO(data), "application/octet-stream")}
    form = {"question_id": question_id} if question_id else {}
    return client.post(f"{API}/response-sets/{set_id}/sections/{code}/evidence", files=files, data=form,
                       headers=auth(user))


def _fetch(client, user, url_response):
    url = url_response.json()["url"]
    return client.get(url, headers=auth(user))


# ===========================================================================
# Evidence
# ===========================================================================

def test_evidence_is_checked_by_content_stored_with_checksum_and_locked_on_submit(client, session, tmp_path):
    _set_storage(tmp_path)
    w = _world(session)
    set_id = _create(client, w["uploader"])

    r = _upload(client, w["uploader"], set_id, "P4", "Stakeholder register.pdf", PDF, "P4.E1")
    assert r.status_code == 201, r.text
    ev = r.json()
    assert ev["sha256"] == hashlib.sha256(PDF).hexdigest() and ev["kind"] == "PDF"
    assert ev["uploaded_by"] == w["uploader"].name and ev["question_id"] == "P4.E1"

    # The extension is not trusted: an executable renamed to .pdf is refused
    r = _upload(client, w["uploader"], set_id, "P4", "invoice.pdf", b"MZ\x90\x00" + b"\x00" * 200)
    assert r.status_code == 400 and r.json()["error"]["code"] == "file_type_not_allowed"
    assert _upload(client, w["uploader"], set_id, "P4", "empty.pdf", b"").status_code == 400
    # ...and the extension must agree with the content (a CSV renamed to .pdf is refused)
    r = _upload(client, w["uploader"], set_id, "P4", "figures.pdf", b"site,kwh\nMumbai,1200\n")
    assert r.status_code == 400 and r.json()["error"]["code"] == "file_type_mismatch"
    assert _upload(client, w["uploader"], set_id, "P4", "figures.csv", b"site,kwh\nMumbai,1200\n").status_code == 201
    r = _upload(client, w["uploader"], set_id, "P4", "big.pdf", PDF + b"0" * (25 * 1024 * 1024))
    assert r.status_code == 400 and r.json()["error"]["code"] == "file_too_large"
    assert _upload(client, w["uploader"], set_id, "P4", "photo.png", PNG).status_code == 201
    r = _upload(client, w["uploader"], set_id, "P4", "x.pdf", PDF, "P6.E1")  # question from another section
    assert r.status_code == 422

    # Download: a short-lived signed URL; the bytes are the bytes uploaded.
    # A non-Latin filename survives the download header.
    r = _upload(client, w["uploader"], set_id, "P4", "बिजली बिल.pdf", PDF)
    hindi = r.json()["public_id"]
    got = _fetch(client, w["uploader"], client.get(
        f"{API}/response-sets/{set_id}/evidence/{hindi}/download", headers=auth(w["uploader"])))
    assert got.status_code == 200 and got.content == PDF
    assert "filename*=UTF-8''" in got.headers["content-disposition"]

    # Another company cannot see or fetch it
    assert client.get(f"{API}/response-sets/{set_id}/evidence", headers=auth(w["outsider"])).status_code == 404

    # Only the uploader can withdraw; withdrawn files disappear from the list but are kept
    r = client.delete(f"{API}/response-sets/{set_id}/evidence/{ev['public_id']}", headers=auth(w["uploader2"]))
    assert r.status_code == 403
    r = client.delete(f"{API}/response-sets/{set_id}/evidence/{ev['public_id']}", headers=auth(w["uploader"]))
    assert r.status_code == 200 and r.json()["withdrawn"] is True
    listed = client.get(f"{API}/response-sets/{set_id}/evidence?section=P4", headers=auth(w["uploader"])).json()
    assert ev["public_id"] not in {e["public_id"] for e in listed} and len(listed) == 3

    # The section detail carries its evidence; once submitted, evidence is locked
    _save(client, w["uploader"], set_id, {q: v for q, v in brsr_fixture.answers().items() if q.startswith("P4.")})
    _approve_answers(client, w["uploader"], w["approver"], set_id, "P4")
    assert client.post(f"{API}/response-sets/{set_id}/sections/P4/submit", headers=auth(w["uploader"])).status_code == 200
    detail = client.get(f"{API}/response-sets/{set_id}/sections/P4", headers=auth(w["uploader"])).json()
    assert len(detail["evidence"]) == 3
    r = _upload(client, w["uploader"], set_id, "P4", "late.pdf", PDF)
    assert r.status_code == 409 and r.json()["error"]["code"] == "section_locked"
    r = client.delete(f"{API}/response-sets/{set_id}/evidence/{hindi}", headers=auth(w["uploader"]))
    assert r.status_code == 409


# ===========================================================================
# Two-tier review
# ===========================================================================

def test_two_tier_review_internal_then_deloitte_by_three_different_people(client, session):
    w = _world(session)
    co = session.get(Company, w["co"].id)
    co.brsr_internal_review_required = True
    session.add(co)
    session.commit()
    set_id = _create(client, w["uploader"])
    _save(client, w["uploader"], set_id, {q: v for q, v in brsr_fixture.answers().items() if q.startswith("P4.")})
    _approve_answers(client, w["uploader"], w["approver"], set_id, "P4")
    assert client.post(f"{API}/response-sets/{set_id}/sections/P4/submit", headers=auth(w["uploader"])).status_code == 200
    base = f"{API}/response-sets/{set_id}/sections/P4"

    # Deloitte cannot jump the queue; it is not even in Deloitte's queue yet
    r = client.post(f"{base}/decide", json={"decision": "approved", "note": "ok"}, headers=auth(w["reviewer"]))
    assert r.status_code == 409 and r.json()["error"]["code"] == "internal_review_pending"
    assert client.get(f"{API}/review-queue", headers=auth(w["reviewer"])).json() == []
    q = client.get(f"{API}/review-queue", headers=auth(w["approver"])).json()
    assert [(x["section"], x["stage"]) for x in q] == [("P4", "internal")]

    r = client.post(f"{base}/decide", json={"decision": "approved", "note": "Checked against the register."},
                    headers=auth(w["approver"]))
    assert r.status_code == 200 and r.json()["status"] == "internally_approved"
    assert r.json()["internal_reviewed_by"] == w["approver"].name and r.json()["editable"] is False
    assert _save(client, w["uploader"], set_id, {"P4.E1": {"text": "x"}}, {"P4.E1": 1}).status_code == 409

    # Now it is Deloitte's, and only Deloitte's
    assert client.get(f"{API}/review-queue", headers=auth(w["approver"])).json() == []
    q = client.get(f"{API}/review-queue", headers=auth(w["reviewer"])).json()
    assert q[0]["stage"] == "final" and q[0]["internal_reviewed_by"] == w["approver"].name
    r = client.post(f"{base}/decide", json={"decision": "approved", "note": "again"}, headers=auth(w["approver"]))
    assert r.status_code == 409 and r.json()["error"]["code"] == "awaiting_deloitte"
    r = client.post(f"{base}/decide", json={"decision": "approved", "note": "Verified."}, headers=auth(w["reviewer"]))
    assert r.status_code == 200 and r.json()["status"] == "approved"
    d = r.json()
    assert d["decided_by"] == w["reviewer"].name and d["internal_reviewed_by"] == w["approver"].name

    # Changes requested at the internal tier go back to the preparer
    _save(client, w["uploader"], set_id, {q: v for q, v in brsr_fixture.answers().items() if q.startswith("P7.")})
    _approve_answers(client, w["uploader"], w["approver"], set_id, "P7")
    client.post(f"{API}/response-sets/{set_id}/sections/P7/submit", headers=auth(w["uploader"]))
    r = client.post(f"{API}/response-sets/{set_id}/sections/P7/decide",
                    json={"decision": "changes_requested", "note": "Add the association list."},
                    headers=auth(w["approver"]))
    assert r.status_code == 200 and r.json()["status"] == "changes_requested" and r.json()["editable"] is True


# ===========================================================================
# Assurance
# ===========================================================================

def _approved_brsr(client, w):
    set_id = _create(client, w["uploader"])
    _fill_all(client, w["uploader"], set_id)
    _submit_all(client, w["uploader"], set_id, w["approver"])
    _approve_all(client, w["reviewer"], set_id)
    return set_id


def _detail(client, user, eid):
    r = client.get(f"{A}/engagements/{eid}", headers=auth(user))
    assert r.status_code == 200, r.text
    return r.json()


def test_assurance_engagement_end_to_end(client, session, tmp_path):
    _set_storage(tmp_path)
    w = _world(session)
    set_id = _approved_brsr(client, w)
    # The preparer attached evidence while preparing (before the sections were approved it was editable)
    # -> attach to a reopened P6 later; here: evidence comes through an information request.

    # --- planning -----------------------------------------------------------
    r = client.post(f"{A}/engagements", json={"response_set_id": set_id, "provider_name": "Assure & Co LLP",
                                              "level": "limited", "scope": ["A", "P3", "P6"]},
                    headers=auth(w["dadmin"]))
    assert r.status_code == 201, r.text
    eid = r.json()["public_id"]
    assert r.json()["status"] == "planning" and [i["section_code"] for i in r.json()["items"]] == ["A", "P3", "P6"]
    r = client.post(f"{A}/engagements", json={"response_set_id": set_id, "provider_name": "Again", "level": "limited"},
                    headers=auth(w["dadmin"]))
    assert r.status_code == 409 and r.json()["error"]["code"] == "engagement_exists"

    # Independence: the client's consultant (who approved every section) cannot join
    session.add(ConsultantAssignment(company_id=w["co"].id, consultant_user_id=w["prac2"].id, is_active=True))
    session.commit()
    cands = {c["name"]: c for c in client.get(f"{A}/engagements/{eid}/candidates", headers=auth(w["dadmin"])).json()}
    assert cands["prac2"]["eligible"] is False and "assigned consultant" in cands["prac2"]["conflicts"][0]
    assert cands["partner"]["eligible"] and cands["partner"]["can_sign"] and not cands["prac"]["can_sign"]
    r = client.post(f"{A}/engagements/{eid}/members", json={"user_id": w["prac2"].id, "role": "practitioner"},
                    headers=auth(w["dadmin"]))
    assert r.status_code == 409 and r.json()["error"]["code"] == "independence_conflict"
    r = client.post(f"{A}/engagements/{eid}/members", json={"user_id": w["prac"].id, "role": "partner"},
                    headers=auth(w["dadmin"]))
    assert r.status_code == 422 and r.json()["error"]["code"] == "not_a_partner"
    for uid, role in ((w["partner"].id, "partner"), (w["eqr"].id, "eqr"), (w["prac"].id, "practitioner")):
        r = client.post(f"{A}/engagements/{eid}/members", json={"user_id": uid, "role": role},
                        headers=auth(w["dadmin"]))
        assert r.status_code == 201, r.text

    # Fieldwork waits for every independence declaration
    r = client.post(f"{A}/engagements/{eid}/start", headers=auth(w["dadmin"]))
    assert r.status_code == 409 and r.json()["error"]["code"] == "independence_not_declared"
    r = client.post(f"{A}/engagements/{eid}/requests", json={"section_code": "P6", "title": "x"},
                    headers=auth(w["prac"]))
    assert r.status_code == 403  # undeclared members cannot act
    for u in ("partner", "eqr", "prac"):
        assert client.post(f"{A}/engagements/{eid}/independence", headers=auth(w[u])).status_code == 200
    # On the team and able to manage it are separate: the partner is both, the practitioner only a member
    assert _detail(client, w["partner"], eid)["can_manage"] is True
    assert _detail(client, w["prac"], eid)["can_manage"] is False
    r = client.post(f"{A}/engagements/{eid}/start", headers=auth(w["dadmin"]))
    assert r.status_code == 200 and r.json()["status"] == "fieldwork"

    # --- information request answered with evidence --------------------------
    r = client.post(f"{A}/engagements/{eid}/requests", json={
        "section_code": "P6", "question_id": "P6.E1", "title": "Electricity bills for March",
        "detail": "All meters, both offices."}, headers=auth(w["prac"]))
    assert r.status_code == 201
    req = r.json()["requests"][0]
    rid = req["public_id"]
    assert req["requested_by"] == w["prac"].name and req["status"] == "open"
    # assurers cannot answer, the client cannot accept
    assert client.post(f"{A}/engagements/{eid}/requests/{rid}/respond", json={"response": "x"},
                       headers=auth(w["prac"])).status_code == 403
    r = client.post(f"{A}/engagements/{eid}/requests/{rid}/respond", json={"response": ""}, headers=auth(w["uploader"]))
    assert r.status_code == 422 and r.json()["error"]["code"] == "response_empty"
    r = client.post(f"{A}/engagements/{eid}/requests/{rid}/evidence",
                    files={"file": ("March bills.pdf", io.BytesIO(PDF), "application/pdf")}, headers=auth(w["uploader"]))
    assert r.status_code == 201 and r.json()["for_assurance_request"] is True
    req_ev = r.json()["public_id"]
    r = client.post(f"{A}/engagements/{eid}/requests/{rid}/respond", json={"response": "Bills attached."},
                    headers=auth(w["uploader"]))
    assert r.status_code == 200
    # Sent back once, answered again, then accepted — the history keeps both rounds
    r = client.post(f"{A}/engagements/{eid}/requests/{rid}/close", json={"decision": "return"},
                    headers=auth(w["prac"]))
    assert r.status_code == 422  # a note is required to send it back
    r = client.post(f"{A}/engagements/{eid}/requests/{rid}/close",
                    json={"decision": "return", "note": "Second office missing."}, headers=auth(w["prac"]))
    assert r.json()["requests"][0]["status"] == "open"
    client.post(f"{A}/engagements/{eid}/requests/{rid}/evidence",
                files={"file": ("Office 2.png", io.BytesIO(PNG), "image/png")}, headers=auth(w["uploader"]))
    client.post(f"{A}/engagements/{eid}/requests/{rid}/respond", json={"response": "Second office added."},
                headers=auth(w["uploader"]))
    r = client.post(f"{A}/engagements/{eid}/requests/{rid}/close", json={"decision": "accept"}, headers=auth(w["prac"]))
    req = r.json()["requests"][0]
    assert req["status"] == "accepted" and req["closed_by"] == w["prac"].name and len(req["history"]) == 1
    assert len(req["evidence"]) == 2
    # Evidence relied on by assurance can no longer be withdrawn
    r = client.delete(f"{API}/response-sets/{set_id}/evidence/{req_ev}", headers=auth(w["uploader"]))
    assert r.status_code == 409 and r.json()["error"]["code"] == "evidence_locked"

    # --- findings: one corrected, one left uncorrected -----------------------
    r = client.post(f"{A}/engagements/{eid}/findings", json={
        "section_code": "P3", "question_id": "P3.E9", "severity": "misstatement",
        "description": "Female reviews overstated by 12 against the HRMS extract."}, headers=auth(w["prac"]))
    f1 = r.json()["findings"][0]["public_id"]
    r = client.post(f"{A}/engagements/{eid}/findings", json={
        "section_code": "P6", "severity": "observation",
        "description": "Grid-electricity split between offices is estimated."}, headers=auth(w["prac"]))
    f2 = r.json()["findings"][1]["public_id"]
    for fid, text in ((f1, "Agreed; corrected from the HRMS extract."), (f2, "Meters to be installed next year.")):
        assert client.post(f"{A}/engagements/{eid}/findings/{fid}/respond", json={"response": text},
                           headers=auth(w["admin"])).status_code == 200
    # "Corrected" is only accepted once the client has actually corrected and re-approved the section
    r = client.post(f"{A}/engagements/{eid}/findings/{f1}/resolve", json={"outcome": "corrected"},
                    headers=auth(w["partner"]))
    assert r.status_code == 409 and r.json()["error"]["code"] == "correction_not_approved"

    # Workpapers: tested by the practitioner, reviewed by someone else
    for code in ("A", "P3", "P6"):
        conclusion = "exceptions" if code in ("P3", "P6") else "no_exceptions"
        r = client.post(f"{A}/engagements/{eid}/items/{code}/prepare", json={
            "procedures": f"Agreed {code} figures to source records on a sample basis.", "conclusion": conclusion},
            headers=auth(w["prac"]))
        assert r.status_code == 200, r.text
    r = client.post(f"{A}/engagements/{eid}/items/A/review", json={"decision": "approve"}, headers=auth(w["prac"]))
    assert r.status_code == 403  # practitioners do not review; and never their own work
    for code in ("A", "P3", "P6"):
        assert client.post(f"{A}/engagements/{eid}/items/{code}/review", json={"decision": "approve"},
                           headers=auth(w["partner"])).status_code == 200

    # The client corrects P3: reopen, fix the answer, have it approved, resubmit, re-approve
    r = client.post(f"{API}/response-sets/{set_id}/sections/P3/reopen", json={"reason": "Assurance finding"},
                    headers=auth(w["uploader"]))
    assert r.status_code == 200
    rev = r.json()["answers"]["P3.E15"]["revision"]
    assert _save(client, w["uploader"], set_id, {"P3.E15": {"text": "No incidences of breach were reported; "
                                                                    "register reconciled."}},
                 {"P3.E15": rev}).status_code == 200
    _approve_answers(client, w["uploader"], w["approver"], set_id, "P3")
    client.post(f"{API}/response-sets/{set_id}/sections/P3/submit", headers=auth(w["uploader"]))
    client.post(f"{API}/response-sets/{set_id}/sections/P3/decide", json={"decision": "approved", "note": "ok"},
                headers=auth(w["reviewer"]))
    r = client.post(f"{A}/engagements/{eid}/findings/{f1}/resolve", json={"outcome": "corrected"},
                    headers=auth(w["partner"]))
    assert r.status_code == 200
    r = client.post(f"{A}/engagements/{eid}/findings/{f2}/resolve", json={"outcome": "uncorrected"},
                    headers=auth(w["partner"]))
    assert r.status_code == 422  # leaving something uncorrected needs a reason
    r = client.post(f"{A}/engagements/{eid}/findings/{f2}/resolve",
                    json={"outcome": "uncorrected", "note": "Not material: under 1% of total energy."},
                    headers=auth(w["partner"]))
    assert r.status_code == 200 and r.json()["counts"]["uncorrected"] == 1

    # P3 changed after it was tested: the reviewed workpaper is now outdated and blocks signing
    d = _detail(client, w["partner"], eid)
    p3 = next(i for i in d["items"] if i["section_code"] == "P3")
    assert p3["outdated"] is True
    assert any("changed after testing" in b for b in d["signoff_blockers"])
    client.post(f"{A}/engagements/{eid}/items/P3/prepare", json={
        "procedures": "Re-performed after correction; agreed to HRMS.", "conclusion": "exceptions"},
        headers=auth(w["prac"]))
    client.post(f"{A}/engagements/{eid}/items/P3/review", json={"decision": "approve"}, headers=auth(w["partner"]))

    # --- quality review and signature ----------------------------------------
    r = client.post(f"{A}/engagements/{eid}/sign", json={"opinion": "unmodified"}, headers=auth(w["partner"]))
    assert r.status_code == 409 and "quality reviewer" in r.json()["error"]["message"]
    assert client.post(f"{A}/engagements/{eid}/sign", json={"opinion": "unmodified"},
                       headers=auth(w["prac"])).status_code == 403
    assert client.post(f"{A}/engagements/{eid}/eqr-clear", json={"note": "Reviewed significant judgements."},
                       headers=auth(w["eqr"])).status_code == 200
    r = client.post(f"{A}/engagements/{eid}/sign", json={"opinion": "qualified"}, headers=auth(w["partner"]))
    assert r.status_code == 422  # a modified opinion needs its basis
    r = client.post(f"{A}/engagements/{eid}/sign", json={"opinion": "unmodified"}, headers=auth(w["partner"]))
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["status"] == "signed" and d["signed_by"] == w["partner"].name and d["reassurance_required"] is False
    assert "nothing has come to our attention" in d["statement_text"]

    # The sign-off register names everyone
    reg = {row["section"]: row for row in d["register"]}
    assert reg["P3"]["prepared_by"] == [w["uploader"].name] and reg["P3"]["approved_by"] == w["reviewer"].name
    assert reg["P3"]["system_sources"] == []
    assert reg["P3"]["tested_by"] == w["prac"].name and reg["P3"]["workpaper_reviewed_by"] == w["partner"].name
    assert reg["P6"]["evidence_files"] == 2

    # Signed documents: stored once, served as signed
    r = client.get(f"{A}/engagements/{eid}/documents/statement", headers=auth(w["admin"]))
    statement = _fetch(client, w["admin"], r).content
    assert hashlib.sha256(statement).hexdigest() == r.json()["sha256"] == d["statement_sha256"]
    text = "\n".join(p.text for p in Document(io.BytesIO(statement)).paragraphs)
    assert "Nuvama Wealth" in text and "Assure & Co LLP" in text and w["partner"].name in text
    r = client.get(f"{A}/engagements/{eid}/documents/register", headers=auth(w["admin"]))
    wb = load_workbook(io.BytesIO(_fetch(client, w["admin"], r).content))
    assert {"Summary", "Sign-off register", "Team & independence", "Information requests", "Findings",
            "Evidence"} <= set(wb.sheetnames)
    team = [row for row in wb["Team & independence"].iter_rows(min_row=2, values_only=True)]
    assert {t[0] for t in team} == {w["partner"].name, w["eqr"].name, w["prac"].name} and all(t[3] for t in team)

    # The pack: documents + evidence by section, every checksum verified
    r = client.get(f"{A}/engagements/{eid}/pack", headers=auth(w["admin"]))
    assert r.status_code == 200
    z = zipfile.ZipFile(io.BytesIO(r.content))
    names = z.namelist()
    assert "MANIFEST.txt" in names and any(n.startswith("evidence/P6/") for n in names)
    manifest = z.read("MANIFEST.txt").decode()
    assert "DOES NOT MATCH" not in manifest and "DIFFERS" not in manifest and "Re-assurance required" in manifest

    # Client view: status and outcomes, but not the assurer's working notes
    cd = _detail(client, w["uploader"], eid)
    assert cd["access"] == "client" and "procedures" not in cd["items"][0]
    assert cd["items"][0]["prepared_by"] == w["prac"].name  # after signing, who tested is disclosed

    # --- the BRSR report carries the assurance -------------------------------
    sec_a = client.get(f"{API}/response-sets/{set_id}/sections/A", headers=auth(w["uploader"])).json()
    assert sec_a["answers"]["A.I"]["value"]["fields"]["assurance_provider"] == "Assure & Co LLP"
    assert sec_a["answers"]["A.I"]["value"]["fields"]["assurance_type"] == "Limited assurance"
    from app.brsr import docx_report, service as brsr_service
    from app.models.brsr import BrsrResponseSet
    rs = session.exec(select(BrsrResponseSet).where(BrsrResponseSet.public_id == set_id)).one()
    snap = brsr_service.frozen_snapshot(session, rs)
    doc = Document(io.BytesIO(docx_report.build_brsr_docx(snap)))
    body = "\n".join(p.text for p in doc.paragraphs)
    cells = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "INDEPENDENT ASSURANCE AND SIGN-OFF" in body and "Assure & Co LLP" in body
    assert w["prac"].name in cells and w["reviewer"].name in cells

    # --- data changes after signing: the assurance goes stale, visibly -------
    client.post(f"{API}/response-sets/{set_id}/sections/P6/reopen", json={"reason": "Typo"},
                headers=auth(w["uploader"]))
    p6 = client.get(f"{API}/response-sets/{set_id}/sections/P6", headers=auth(w["uploader"])).json()
    qid = next(q for q, a in p6["answers"].items() if a["value"] and "text" in (a["value"] or {}))
    assert _save(client, w["uploader"], set_id, {qid: {"text": "Edited after assurance."}},
                 {qid: p6["answers"][qid]["revision"]}).status_code == 200
    d = _detail(client, w["admin"], eid)
    assert d["reassurance_required"] is True
    # The register stays exactly as signed, even though P6 has since been reopened
    assert d["register_frozen"] is True
    row = {r["section"]: r for r in d["register"]}["P6"]
    assert row["section_status"] == "approved" and row["approved_by"] == w["reviewer"].name
    sec_a = client.get(f"{API}/response-sets/{set_id}/sections/A", headers=auth(w["uploader"])).json()
    assert sec_a["answers"]["A.I"]["value"]["fields"].get("assurance_provider") != "Assure & Co LLP"

    # Re-assurance: a new engagement supersedes the stale one, which stays on record
    sets = client.get(f"{A}/assurable-sets", headers=auth(w["dadmin"])).json()
    assert sets[0]["engagement"]["needs_reassurance"] is True
    r = client.post(f"{A}/engagements", json={"response_set_id": set_id, "provider_name": "Assure & Co LLP",
                                              "level": "limited", "scope": ["P6"]}, headers=auth(w["dadmin"]))
    assert r.status_code == 201, r.text
    old = _detail(client, w["admin"], eid)
    assert old["status"] == "superseded" and old["superseded_by"] == r.json()["public_id"]
    assert old["statement_sha256"] == d["statement_sha256"]  # the signed record is untouched


def test_assurance_access_is_need_to_know(client, session):
    w = _world(session)
    set_id = _create(client, w["uploader"])
    r = client.post(f"{A}/engagements", json={"response_set_id": set_id, "provider_name": "Assure & Co LLP"},
                    headers=auth(w["partner"]))
    assert r.status_code == 201  # an assurance partner can open an engagement in their organisation
    eid = r.json()["public_id"]
    sets = client.get(f"{A}/assurable-sets", headers=auth(w["partner"])).json()
    assert [(x["company_name"], x["engagement"]["public_id"]) for x in sets] == [("Nuvama Wealth", eid)]
    assert "answers" not in sets[0]
    assert client.get(f"{A}/assurable-sets", headers=auth(w["prac"])).status_code == 403
    # The client sees its own engagement; another company does not know it exists
    assert _detail(client, w["admin"], eid)["access"] == "client"
    assert client.get(f"{A}/engagements/{eid}", headers=auth(w["outsider"])).status_code == 404
    assert client.get(f"{A}/engagements", headers=auth(w["outsider"])).json() == []
    # An assurance practitioner not on the team sees nothing, including the client's BRSR
    assert client.get(f"{A}/engagements/{eid}", headers=auth(w["prac"])).status_code == 404
    assert client.get(f"{API}/response-sets/{set_id}", headers=auth(w["prac"])).status_code == 404
    # The assigned consultant can follow it, read-only
    assert _detail(client, w["reviewer"], eid)["access"] == "observer"
    assert client.post(f"{A}/engagements/{eid}/members", json={"user_id": w["prac"].id, "role": "practitioner"},
                       headers=auth(w["reviewer"])).status_code == 403
    # Clients cannot create engagements or act as assurers
    assert client.post(f"{A}/engagements", json={"response_set_id": set_id, "provider_name": "Me"},
                       headers=auth(w["admin"])).status_code == 403
    # Once on the team, the practitioner reads the BRSR through the engagement only, and read-only
    client.post(f"{A}/engagements/{eid}/members", json={"user_id": w["prac"].id, "role": "practitioner"},
                headers=auth(w["partner"]))
    s = client.get(f"{A}/engagements/{eid}/sections/P4", headers=auth(w["prac"]))
    assert s.status_code == 200 and s.json()["editable"] is False
    assert client.get(f"{A}/engagements/{eid}/sections/P4", headers=auth(w["admin"])).status_code == 404
    # Withdrawal frees the BRSR for a new engagement
    r = client.post(f"{A}/engagements/{eid}/withdraw", json={"reason": "Scope changed"}, headers=auth(w["partner"]))
    assert r.json()["status"] == "withdrawn"
    assert client.post(f"{A}/engagements", json={"response_set_id": set_id, "provider_name": "Assure & Co LLP"},
                       headers=auth(w["partner"])).status_code == 201


def test_statement_wording_follows_level_and_opinion():
    from app.assurance.documents import conclusion_paragraph
    from app.models.assurance import AssuranceEngagement

    def text(level, opinion):
        return conclusion_paragraph(AssuranceEngagement(level=level, opinion=opinion, financial_year="2024-25",
                                                        provider_name="x", company_id=1, response_set_id=1,
                                                        created_by=1), "ACME")
    assert text("limited", "unmodified").startswith("Based on the procedures")
    assert text("reasonable", "unmodified").startswith("In our opinion")
    assert text("limited", "qualified").startswith("Except for")
    assert "is not prepared" in text("reasonable", "adverse")
    assert "do not express a conclusion" in text("limited", "disclaimer")


def test_migration_rbac_data_matches_seed_definitions():
    """Production gets roles from the migration, dev/test from the seeder: they must agree."""
    import importlib.util
    import pathlib
    from app.rbac.definitions import PERMISSIONS, ROLE_PERMISSIONS, ROLES
    path = next(pathlib.Path("alembic/versions").glob("e6b2d9a4c7f1_*.py"))
    spec = importlib.util.spec_from_file_location("mig", path)
    mig = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mig)
    perms = dict(PERMISSIONS)
    assert all(perms.get(c) == d for c, d in mig.NEW_PERMISSIONS)
    roles = {c: (s, d) for c, s, d in ROLES}
    assert all(roles.get(c) == (s, d) for c, s, d in mig.NEW_ROLES)
    for role, granted in ROLE_PERMISSIONS.items():
        seeded = {p for p in granted if p.startswith("assurance:")}
        migrated = {p for p in mig.GRANTS.get(role, []) if p.startswith("assurance:")}
        assert seeded == migrated, role
    for role in ("Assurance Partner", "Assurance Practitioner"):
        assert set(mig.GRANTS[role]) == set(ROLE_PERMISSIONS[role]), role

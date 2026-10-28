"""Surveys & assessments: builder rules, scoring, lifecycle, audience,
anonymity, links, files, results and tenancy."""
import io

import pytest
from openpyxl import load_workbook
from sqlmodel import select

from app.core import ratelimit
from app.models.audit_log import AuditLog
from app.models.notification import NotificationOutbox
from app.models.survey import SurveyFile, SurveyInvite, SurveyResponse
from app.surveys import schema
from tests.conftest_helpers import auth, bootstrap, make_company, make_user
from tests.test_layer1_kpi import _set_storage

API = "/api/v1"


@pytest.fixture(autouse=True)
def _fresh_limits():
    ratelimit.reset()
    yield
    ratelimit.reset()


def _world(session):
    org = bootstrap(session)
    co = make_company(session, org, "Acme Ltd")
    admin = make_user(session, "admin@acme.com", "client", "Client Administrator", company=co)
    up = make_user(session, "up@acme.com", "client", "Client Uploader", company=co)
    appr = make_user(session, "appr@acme.com", "client", "Client Approver", company=co)
    up.department, appr.department = "HR", "HR"
    session.add(up)
    session.add(appr)
    session.commit()
    return org, co, admin, up, appr


DEF = {
    "sections": [
        {"id": "s1", "title": "Basics", "questions": [
            {"id": "q1", "type": "single_choice", "title": "Policy?", "required": True,
             "options": [{"id": "y", "label": "Yes", "score": 4}, {"id": "n", "label": "No", "score": 0},
                         {"id": "na", "label": "Not applicable", "na": True}]},
            {"id": "q2", "type": "long_text", "title": "Why not?", "required": True,
             "show_if": {"match": "all", "rules": [{"question": "q1", "op": "equals", "value": "n"}]}},
            {"id": "q3", "type": "multi_choice", "title": "Tracked", "allow_other": True,
             "options": [{"id": "a", "label": "Energy", "score": 1}, {"id": "b", "label": "Water", "score": 1},
                         {"id": "none", "label": "None", "exclusive": True}]},
        ]},
        {"id": "s2", "title": "People", "weight": 3, "questions": [
            {"id": "q4", "type": "scale", "title": "Safety", "min": 1, "max": 5, "required": True},
            {"id": "q5", "type": "matrix", "title": "Practices", "rows": [{"id": "r1", "label": "Training"},
                                                                          {"id": "r2", "label": "Benefits"}],
             "columns": [{"id": "c0", "label": "No", "score": 0}, {"id": "c1", "label": "Yes", "score": 2},
                         {"id": "cna", "label": "N/A", "na": True}]},
            {"id": "q6", "type": "nps", "title": "Recommend?"},
            {"id": "q7", "type": "number", "title": "Headcount", "min": 0, "integer": True},
            {"id": "q8", "type": "file", "title": "Policy document", "max_files": 2},
        ]},
    ],
    "bands": [{"min": 50, "label": "Good"}, {"min": 0, "label": "Low"}],
}


# ---------------------------------------------------------------------------
# Pure rules
# ---------------------------------------------------------------------------

def test_definition_is_checked_and_normalised():
    d = schema.normalize_definition(DEF)
    assert [b["label"] for b in d["bands"]] == ["Low", "Good"]
    assert d["sections"][1]["questions"][0]["max"] == 5
    bad = [
        {"sections": [{"id": "s", "questions": [{"id": "q", "type": "telepathy"}]}]},
        {"sections": [{"id": "s", "questions": [{"id": "q", "type": "short_text"}, {"id": "q", "type": "date"}]}]},
        # a rule may only look back
        {"sections": [{"id": "s", "questions": [
            {"id": "a", "type": "short_text", "show_if": {"rules": [{"question": "b", "op": "answered"}]}},
            {"id": "b", "type": "short_text"}]}]},
        # an operator that makes no sense for the type
        {"sections": [{"id": "s", "questions": [
            {"id": "a", "type": "short_text"},
            {"id": "b", "type": "short_text", "show_if": {"rules": [{"question": "a", "op": "gte", "value": 1}]}}]}]},
        {"sections": [{"id": "s", "questions": [{"id": "a", "type": "scale", "min": 2, "max": 5}]}]},
        {"sections": [], "bands": [{"min": 10, "label": "A"}, {"min": 10, "label": "B"}]},
    ]
    for b in bad:
        with pytest.raises(schema.DefinitionError):
            schema.normalize_definition(b)


def test_publish_problems_and_structure_lock():
    blank = schema.normalize_definition({"sections": [{"id": "s", "title": "", "questions": [
        {"id": "q", "type": "single_choice", "title": "", "options": [{"id": "o", "label": ""}]}]}]})
    msgs = " ".join(p["message"] for p in schema.publish_problems(blank, "assessment"))
    assert "has no text" in msgs and "at least 2 options" in msgs and "needs scores" in msgs

    d = schema.normalize_definition(DEF)
    reworded = schema.normalize_definition({**DEF, "sections": [
        {**DEF["sections"][0], "questions": [{**DEF["sections"][0]["questions"][0], "title": "Policy in place?",
                                              "options": DEF["sections"][0]["questions"][0]["options"]
                                              + [{"id": "new", "label": "Partly", "score": 2}]}]
         + DEF["sections"][0]["questions"][1:]}, DEF["sections"][1]]})
    assert schema.locked_changes(d, reworded) == []
    removed = schema.normalize_definition({**DEF, "sections": [DEF["sections"][0]]})
    assert any("was removed" in p for p in schema.locked_changes(d, removed))


def test_logic_answers_and_scoring():
    d = schema.normalize_definition(DEF)
    # Hidden question: not required, and its stale answer is dropped
    clean = schema.clean_answers(d, {"q1": "y", "q2": "left over", "q4": 3}, final=True)
    assert "q2" not in clean
    with pytest.raises(schema.AnswerError) as e:
        schema.clean_answers(d, {"q1": "n", "q4": 3}, final=True)
    assert e.value.details["errors"] == {"q2": "This question needs an answer"}
    # Drafts don't enforce required
    assert schema.clean_answers(d, {"q1": "n"}, final=False) == {"q1": "n"}
    errs = None
    try:
        schema.clean_answers(d, {"q1": "zzz", "q3": ["a", "none"], "q4": 9, "q7": 1.5,
                                 "q5": {"r1": "nope"}}, final=False)
    except schema.AnswerError as ex:
        errs = ex.details["errors"]
    assert set(errs) == {"q1", "q3", "q4", "q7", "q5"}
    assert "can't be combined" in errs["q3"]
    with pytest.raises(schema.AnswerError):  # "Other" needs its text
        schema.clean_answers(d, {"q3": ["__other__"]}, final=False)
    assert schema.clean_answers(d, {"q3": ["__other__"], "q3.other": "Waste"}, final=False)["q3.other"] == "Waste"

    # Scoring: s1 = q1 (4/4) + q3 (1/2) → 75 %; s2 = q4 (2/4 = 50 %) + q5 (one row N/A, other 'Yes' → 100 %) → 75 %
    a = {"q1": "y", "q3": ["a"], "q4": 3, "q5": {"r1": "c1", "r2": "cna"}}
    s = schema.score_response(d, schema.clean_answers(d, a, final=True))
    assert s["sections"]["s1"]["percent"] == 75.0
    assert s["sections"]["s2"]["percent"] == 75.0
    assert s["percent"] == 75.0 and s["band"] == "Good"
    # Not applicable takes the question out; unanswered visible scored questions count as zero
    s = schema.score_response(d, {"q1": "na", "q4": 1})
    assert s["questions"]["q1"]["na"] and s["sections"]["s1"]["percent"] == 0.0  # q3 unanswered → 0
    assert s["percent"] == 0.0 and s["band"] == "Low"
    # Section weights: s2 counts three times as much as s1
    s = schema.score_response(d, {"q1": "y", "q3": ["a", "b"], "q4": 1, "q5": {"r1": "c0", "r2": "c0"}})
    assert s["percent"] == 25.0
    # Respondents never see scores
    view = schema.respondent_view(d)
    assert all("score" not in o for o in view["sections"][0]["questions"][0]["options"])


# ---------------------------------------------------------------------------
# API
# ---------------------------------------------------------------------------

def _create(client, user, **body):
    r = client.post(f"{API}/surveys", json=body, headers=auth(user))
    assert r.status_code == 201, r.text
    return r.json()


def _publishable(client, user, kind="assessment", **settings):
    sv = _create(client, user, title="Supplier check", kind=kind)
    r = client.put(f"{API}/surveys/{sv['id']}", json={"revision": sv["revision"], "definition": DEF, **settings},
                   headers=auth(user))
    assert r.status_code == 200, r.text
    return r.json()


def test_templates_and_builder_lifecycle(client, session):
    _, co, admin, up, _ = _world(session)
    h = auth(admin)
    tpls = client.get(f"{API}/surveys/templates", headers=h).json()
    assert {t["id"] for t in tpls if t["built_in"]} >= {"tpl_esg_maturity", "tpl_supplier_dd", "tpl_climate_risk",
                                                        "tpl_materiality", "tpl_wellbeing", "tpl_customer"}
    full = client.get(f"{API}/surveys/templates/tpl_supplier_dd", headers=h).json()
    assert full["definition"]["bands"][0]["label"] == "High risk"

    sv = _create(client, admin, template_id="tpl_esg_maturity")
    assert sv["kind"] == "assessment" and sv["status"] == "draft" and sv["problems"] == []
    # Optimistic concurrency
    ok = client.put(f"{API}/surveys/{sv['id']}", json={"revision": sv["revision"], "title": "FY25 baseline"}, headers=h)
    assert ok.status_code == 200 and ok.json()["revision"] == sv["revision"] + 1
    stale = client.put(f"{API}/surveys/{sv['id']}", json={"revision": sv["revision"], "title": "x"}, headers=h)
    assert stale.status_code == 409 and stale.json()["error"]["code"] == "survey_conflict"

    blank = _create(client, admin, title="Empty")
    r = client.post(f"{API}/surveys/{blank['id']}/publish", headers=h)
    assert r.status_code == 422 and r.json()["error"]["code"] == "survey_not_ready"
    assert r.json()["error"]["details"]["problems"]
    # A draft with no responses can be deleted
    assert client.delete(f"{API}/surveys/{blank['id']}", headers=h).status_code == 204

    # Save as a company template, then build from it
    t = client.post(f"{API}/surveys/{sv['id']}/save-as-template", json={"title": "Our maturity check"}, headers=h).json()
    assert t["is_template"] and t["status"] == "template"
    assert any(x["id"] == t["id"] and not x["built_in"] for x in client.get(f"{API}/surveys/templates", headers=h).json())
    from_company = _create(client, admin, template_id=t["id"])
    assert from_company["definition"] == t["definition"]
    assert client.post(f"{API}/surveys/{t['id']}/publish", headers=h).status_code == 409

    dup = client.post(f"{API}/surveys/{sv['id']}/duplicate", headers=h).json()
    assert dup["title"].startswith("Copy of") and dup["status"] == "draft"
    listed = client.get(f"{API}/surveys", headers=h).json()
    assert {x["id"] for x in listed} >= {sv["id"], dup["id"], from_company["id"]} and t["id"] not in {x["id"] for x in listed}


def test_internal_audience_respond_and_results(client, session):
    _, co, admin, up, appr = _world(session)
    h = auth(admin)
    sv = _publishable(client, admin)
    sid = sv["id"]
    r = client.post(f"{API}/surveys/{sid}/invites", json={"departments": ["hr"], "user_ids": [admin.id]}, headers=h)
    assert r.status_code == 201 and r.json()["added"] == 3
    assert client.post(f"{API}/surveys/{sid}/invites", json={"user_ids": [up.id]}, headers=h).json()["skipped"] == 1
    # Not yet published: invitees can't see it
    assert client.get(f"{API}/my-surveys", headers=auth(up)).json() == []

    assert client.post(f"{API}/surveys/{sid}/publish", headers=h).status_code == 200
    out = session.exec(select(NotificationOutbox).where(NotificationOutbox.event_type == "survey_invited")).all()
    assert sorted(out[-1].payload["recipient_user_ids"]) == sorted([admin.id, up.id, appr.id])

    mine = client.get(f"{API}/my-surveys", headers=auth(up)).json()
    assert [m["id"] for m in mine] == [sid] and mine[0]["my_status"] == "not_started"
    form = client.get(f"{API}/my-surveys/{sid}", headers=auth(up)).json()
    assert "score" not in form["survey"]["definition"]["sections"][0]["questions"][0]["options"][0]

    # Save a draft, come back, submit
    d = client.put(f"{API}/my-surveys/{sid}/response", json={"answers": {"q1": "n"}}, headers=auth(up)).json()
    assert d["status"] == "in_progress" and d["progress"]["total"] == 8  # q2 now shown
    bad = client.post(f"{API}/my-surveys/{sid}/submit", json={"answers": {"q1": "n"}}, headers=auth(up))
    assert bad.status_code == 422 and set(bad.json()["error"]["details"]["errors"]) == {"q2", "q4"}
    conflict = client.put(f"{API}/my-surveys/{sid}/response", json={"answers": {}, "revision": 999}, headers=auth(up))
    assert conflict.status_code == 409
    ok = client.post(f"{API}/my-surveys/{sid}/submit", json={"answers": {
        "q1": "n", "q2": "Not yet drafted", "q3": ["a"], "q4": 2, "q5": {"r1": "c1", "r2": "c0"}, "q6": 9,
        "q7": 120}}, headers=auth(up))
    assert ok.status_code == 200, ok.text
    assert client.post(f"{API}/my-surveys/{sid}/submit", json={"answers": {}}, headers=auth(up)).status_code == 409
    client.post(f"{API}/my-surveys/{sid}/submit", json={"answers": {
        "q1": "y", "q3": ["a", "b"], "q4": 5, "q5": {"r1": "c1", "r2": "c1"}, "q6": 3}}, headers=auth(appr))

    res = client.get(f"{API}/surveys/{sid}/results", headers=h).json()
    assert res["summary"]["responses"] == 2 and res["summary"]["response_rate"] == 67
    q1 = next(q for q in res["questions"] if q["id"] == "q1")
    assert {o["id"]: o["count"] for o in q1["options"]} == {"y": 1, "n": 1, "na": 0}
    q6 = next(q for q in res["questions"] if q["id"] == "q6")
    assert q6["nps"] == {"score": 0, "promoters": 50.0, "passives": 0.0, "detractors": 50.0, "n": 2}
    q2 = next(q for q in res["questions"] if q["id"] == "q2")
    assert q2["shown"] == 1 and q2["answers"][0]["text"] == "Not yet drafted"
    assert res["scores"]["average"] is not None and res["findings"]
    assert {g["label"] for g in res["groups"]} == {"HR"}

    lst = client.get(f"{API}/surveys/{sid}/responses", headers=h).json()
    assert len(lst) == 2 and {x["respondent"] for x in lst} == {"up", "appr"}
    det = client.get(f"{API}/surveys/{sid}/responses/{lst[-1]['id']}", headers=h).json()
    assert det["sections"][0]["questions"][0]["value"] == "No"

    x = client.get(f"{API}/surveys/{sid}/export", headers=h)
    assert x.status_code == 200
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Summary", "Responses", "Questions"] and wb["Responses"].max_row == 3

    # Structure lock: rewording is fine, removing a question is not; score changes re-score
    cur = client.get(f"{API}/surveys/{sid}", headers=h).json()
    assert cur["structure_locked"]
    gone = {**cur["definition"], "sections": [cur["definition"]["sections"][0]]}
    r = client.put(f"{API}/surveys/{sid}", json={"revision": cur["revision"], "definition": gone}, headers=h)
    assert r.status_code == 409 and r.json()["error"]["code"] == "survey_structure_locked"
    rescored = cur["definition"]
    rescored["sections"][0]["questions"][0]["options"][1]["score"] = 4  # "No" now earns full marks
    r = client.put(f"{API}/surveys/{sid}", json={"revision": cur["revision"], "definition": rescored}, headers=h)
    assert r.status_code == 200
    resp = session.exec(select(SurveyResponse).where(SurveyResponse.survey_id == session.exec(
        select(SurveyInvite.survey_id).where(SurveyInvite.user_id == up.id)).first())).all()
    assert all(x.score["questions"]["q1"]["points"] == 4 for x in resp)
    assert client.put(f"{API}/surveys/{sid}", json={"anonymous": True}, headers=h).status_code == 409

    # Remind, close, reopen; closed surveys refuse answers; with responses: archive, not delete
    assert client.post(f"{API}/surveys/{sid}/remind", headers=h).json()["reminded"] == 1  # only admin is pending
    assert client.post(f"{API}/surveys/{sid}/close", headers=h).json()["status"] == "closed"
    r = client.put(f"{API}/my-surveys/{sid}/response", json={"answers": {"q1": "y"}}, headers=auth(admin))
    assert r.status_code == 409 and r.json()["error"]["code"] == "survey_closed"
    assert client.post(f"{API}/surveys/{sid}/reopen", headers=h).json()["status"] == "open"
    assert client.delete(f"{API}/surveys/{sid}", headers=h).status_code == 409
    assert client.post(f"{API}/surveys/{sid}/archive", headers=h).json()["status"] == "archived"
    assert client.get(f"{API}/my-surveys/{sid}", headers=auth(admin)).status_code == 404


def test_open_to_company_and_anonymity(client, session):
    org, co, admin, up, appr = _world(session)
    h = auth(admin)
    extra = [make_user(session, f"e{i}@acme.com", "client", "Client Uploader", company=co) for i in range(4)]
    for u in extra:
        u.department = "Ops"
        session.add(u)
    session.commit()
    sv = _publishable(client, admin, kind="survey", anonymous=True, open_to_company=True)
    sid = sv["id"]
    assert client.post(f"{API}/surveys/{sid}/publish", headers=h).status_code == 200
    ans = {"q1": "y", "q4": 4, "q6": 10}
    people = [up, appr] + extra
    for u in people[:4]:
        assert client.post(f"{API}/my-surveys/{sid}/submit", json={"answers": ans}, headers=auth(u)).status_code == 200
    # Below the threshold: nothing shown; individual responses never
    res = client.get(f"{API}/surveys/{sid}/results", headers=h).json()
    assert res["withheld"] and res["questions"] == []
    assert client.get(f"{API}/surveys/{sid}/responses", headers=h).status_code == 403
    assert client.get(f"{API}/surveys/{sid}/export", headers=h).status_code == 409
    for u in people[4:]:
        client.post(f"{API}/my-surveys/{sid}/submit", json={"answers": ans}, headers=auth(u))
    res = client.get(f"{API}/surveys/{sid}/results", headers=h).json()
    assert not res["withheld"] and res["summary"]["responses"] == 6
    # HR has 2 responses → hidden; Ops has 4 → hidden as well (threshold 5)
    assert all(g["hidden"] for g in res["groups"])
    assert client.get(f"{API}/surveys/{sid}/results", params={"group": "Ops"}, headers=h).status_code == 422

    # Nothing links a person to their answers
    rows = session.exec(select(SurveyResponse).where(SurveyResponse.status == "submitted")).all()
    assert all(r.user_id is None and r.invite_id is None and r.respondent_name is None for r in rows)
    assert all(r.submitted_at.hour == 0 and r.submitted_at.minute == 0 for r in rows)
    invs = session.exec(select(SurveyInvite)).all()
    assert len(invs) == 6 and all(i.status == "submitted" and i.submitted_at is None for i in invs)
    logs = session.exec(select(AuditLog).where(AuditLog.action == "survey.response_submitted")).all()
    assert logs and all(lg.actor_user_id is None and lg.entity_type == "survey" for lg in logs)
    # The person still sees that they've answered (but not their answers)
    form = client.get(f"{API}/my-surveys/{sid}", headers=auth(up)).json()
    assert form["response"]["status"] == "submitted" and form["response"]["answers"] == {}
    assert client.post(f"{API}/my-surveys/{sid}/submit", json={"answers": ans}, headers=auth(up)).status_code == 409


def test_links_files_and_rate_limit(client, session, tmp_path):
    _set_storage(tmp_path)
    _, co, admin, *_ = _world(session)
    h = auth(admin)
    sv = _publishable(client, admin, show_score_to_respondent=True)
    sid = sv["id"]
    r = client.post(f"{API}/surveys/{sid}/invites", json={"emails": [
        {"email": "Supplier@Vendor.com", "name": "Vendor Co", "group": "Packaging"}, {"email": "b@vendor.com"}]},
        headers=h)
    links = {x["email"]: x["token"] for x in r.json()["links"]}
    assert set(links) == {"supplier@vendor.com", "b@vendor.com"}
    bad = client.post(f"{API}/surveys/{sid}/invites", json={"emails": [{"email": "not-an-email"}]}, headers=h)
    assert bad.status_code == 422 and bad.json()["error"]["code"] == "invalid_email"
    assert client.get(f"{API}/public/surveys/{links['b@vendor.com']}").status_code == 404  # still a draft

    client.post(f"{API}/surveys/{sid}/publish", headers=h)
    tok = links["supplier@vendor.com"]
    form = client.get(f"{API}/public/surveys/{tok}").json()
    assert form["survey"]["title"] == "Supplier check" and form["respondent"]["name"] == "Vendor Co"
    assert form["survey"]["company"] == "Acme Ltd"

    # Upload a file, then submit referencing it
    pdf = b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF"
    up = client.post(f"{API}/public/surveys/{tok}/files", data={"question_id": "q8"},
                     files={"file": ("policy.pdf", pdf, "application/pdf")})
    assert up.status_code == 201, up.text
    fid = up.json()["file"]["id"]
    assert up.json()["response"]["answers"]["q8"] == [fid]
    wrong = client.post(f"{API}/public/surveys/{tok}/files", data={"question_id": "q1"},
                        files={"file": ("policy.pdf", pdf, "application/pdf")})
    assert wrong.status_code == 422
    sub = client.post(f"{API}/public/surveys/{tok}/submit", json={"answers": {
        "q1": "y", "q3": ["a", "b"], "q4": 5, "q5": {"r1": "c1", "r2": "c1"}, "q8": [fid]}})
    assert sub.status_code == 200 and sub.json()["score"]["percent"] == 100.0 and sub.json()["score"]["band"] == "Good"
    assert client.post(f"{API}/public/surveys/{tok}/submit", json={"answers": {}}).status_code == 409

    # Manager downloads the file through a signed URL
    dl = client.get(f"{API}/surveys/{sid}/files/{fid}/download", headers=h).json()
    got = client.get(dl["url"].replace("http://testserver", ""))
    assert got.status_code == 200 and got.content == pdf

    # Renewing a link stops the old one; the draft stays with the invitation
    b_tok = links["b@vendor.com"]
    client.put(f"{API}/public/surveys/{b_tok}/response", json={"answers": {"q1": "n"}})
    inv_b = next(i for i in client.get(f"{API}/surveys/{sid}/audience", headers=h).json()["invites"]
                 if i["email"] == "b@vendor.com")
    assert inv_b["status"] == "started"
    new = client.post(f"{API}/surveys/{sid}/links", json={"invite_ids": [inv_b["id"]]}, headers=h).json()["links"][0]
    assert client.get(f"{API}/public/surveys/{b_tok}").status_code == 404
    assert client.get(f"{API}/public/surveys/{new['token']}").json()["response"]["answers"] == {"q1": "n"}

    # Public link: anyone, resumable by key, optional name and organisation
    sv2 = client.put(f"{API}/surveys/{sid}/public-link", json={"enabled": True}, headers=h).json()
    pt = sv2["public_token"]
    s1 = client.put(f"{API}/public/surveys/{pt}/response", json={"answers": {"q1": "y"}}).json()
    key = s1["resume_key"]
    assert key and client.get(f"{API}/public/surveys/{pt}").json()["response"]["status"] == "not_started"
    resumed = client.get(f"{API}/public/surveys/{pt}", headers={"X-Resume-Key": key}).json()
    assert resumed["response"]["answers"] == {"q1": "y"} and resumed["survey"]["ask_identity"]
    done = client.post(f"{API}/public/surveys/{pt}/submit", headers={"X-Resume-Key": key}, json={
        "answers": {"q1": "y", "q4": 3}, "name": "Visitor", "organisation": "Logistics Co"})
    assert done.status_code == 200
    rows = client.get(f"{API}/surveys/{sid}/responses", headers=h).json()
    assert {(x["respondent"], x["group"], x["channel"]) for x in rows} >= {("Vendor Co", "Packaging", "email"),
                                                                            ("Visitor", "Logistics Co", "public")}
    # Switching the link off stops it
    client.put(f"{API}/surveys/{sid}/public-link", json={"enabled": False}, headers=h)
    assert client.get(f"{API}/public/surveys/{pt}").status_code == 404

    # Rate limit on submissions from one address
    codes = [client.post(f"{API}/public/surveys/{tok}/submit", json={"answers": {}}).status_code for _ in range(12)]
    assert 429 in codes
    assert session.exec(select(SurveyFile)).first().sha256_checksum


def test_tenancy_and_permissions(client, session):
    org, co, admin, up, appr = _world(session)
    other = make_company(session, org, "Other Co")
    other_admin = make_user(session, "admin@other.com", "client", "Client Administrator", company=other)
    stranger = make_user(session, "s@other.com", "client", "Client Uploader", company=other)
    sv = _publishable(client, admin)
    sid = sv["id"]
    client.post(f"{API}/surveys/{sid}/publish", headers=auth(admin))
    assert client.get(f"{API}/surveys/{sid}", headers=auth(other_admin)).status_code == 404
    assert client.get(f"{API}/surveys", headers=auth(appr)).status_code == 403  # respond-only role
    assert client.get(f"{API}/my-surveys/{sid}", headers=auth(up)).status_code == 404  # not invited
    assert client.get(f"{API}/my-surveys/{sid}", headers=auth(stranger)).status_code == 404
    r = client.post(f"{API}/surveys/{sid}/invites", json={"user_ids": [stranger.id]}, headers=auth(admin))
    assert r.status_code == 422  # only this company's users

    # Deloitte: the assigned consultant manages the client's surveys; others don't see them
    from app.models.consultant_assignment import ConsultantAssignment
    cons = make_user(session, "c@deloitte.com", "deloitte", "Consultant", org=org)
    other_cons = make_user(session, "c2@deloitte.com", "deloitte", "Consultant", org=org)
    session.add(ConsultantAssignment(consultant_user_id=cons.id, company_id=co.id))
    session.commit()
    assert client.get(f"{API}/surveys", params={"company_id": co.id}, headers=auth(cons)).status_code == 200
    assert client.get(f"{API}/surveys/{sid}/results", headers=auth(cons)).status_code == 200
    assert client.get(f"{API}/surveys/{sid}", headers=auth(other_cons)).status_code == 404
    made = client.post(f"{API}/surveys", json={"company_id": co.id, "template_id": "tpl_customer"},
                       headers=auth(cons))
    assert made.status_code == 201 and made.json()["company_id"] == co.id

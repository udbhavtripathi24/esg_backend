"""Data sheets: the engagement's 27 templates as Upload Center types.

Template download → dry-run check → strict upload (row-level errors) →
review → extraction into kpi_values → BRSR figures and coverage.
"""
import io
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook
from sqlmodel import select

from app.datasheets.catalog import IDENTITY, SHEETS
from app.datasheets.seed import ensure_datasheet_types
from app.datasheets.service import parse_sheet, template_workbook
from app.models.kpi import KpiValue
from tests.conftest_helpers import auth, bootstrap, make_company, make_user
from tests.test_layer1_kpi import _approve, _run_kpi_job, _set_storage

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ENERGY = "Nonrenewable_Purchased_Electricity_from_Grid_kWh"


def _book(sheet: str, rows: list[dict]) -> bytes:
    cols = IDENTITY + [c["name"] for c in SHEETS[sheet]["columns"]]
    wb = Workbook()
    ws = wb.active
    ws.append(cols)
    for r in rows:
        ws.append([r.get(c) for c in cols])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _month(y, m, **vals):
    end = date(y, m + 1, 1) if m < 12 else date(y + 1, 1, 1)
    from datetime import timedelta
    return {"Businesses": "Nuvama Wealth", "Start_Date": date(y, m, 1), "End_Date": end - timedelta(days=1),
            "Location": "Mumbai", **vals}


def _world(session):
    org = bootstrap(session)
    ensure_datasheet_types(session)
    co = make_company(session, org, "Nuvama Wealth")
    up = make_user(session, "up@client.com", "client", "Client Uploader", company=co)
    appr = make_user(session, "appr@client.com", "client", "Client Approver", company=co)
    return co, up, appr


def _dataset(client, co, user, sheet):
    r = client.post("/api/v1/datasets", json={"company_id": co.id, "upload_type_code": f"sheet_{sheet}",
                                              "reporting_period_start": "2024-04-01",
                                              "reporting_period_end": "2025-03-31",
                                              "reporting_frequency": "monthly"}, headers=auth(user))
    assert r.status_code == 201, r.text
    ds = r.json()["public_id"]
    v = client.get(f"/api/v1/datasets/{ds}/versions", headers=auth(user)).json()[0]["public_id"]
    return ds, v


def _put(client, user, ds, v, name, data):
    return client.post(f"/api/v1/datasets/{ds}/versions/{v}/files", files={"file": (name, data, XLSX)},
                       data={"role": "data"}, headers=auth(user))


def test_catalogue_and_seed_are_consistent(session):
    assert len(SHEETS) == 27
    first = ensure_datasheet_types(session)
    assert first["upload_types_added"] == 27 and first["kpi_definitions_added"] == 398
    assert ensure_datasheet_types(session) == {"upload_types_added": 0, "kpi_definitions_added": 0}


def test_template_has_the_exact_headers_drop_downs_and_instructions():
    wb = load_workbook(io.BytesIO(template_workbook("diversity")))
    ws = wb.worksheets[0]
    headers = [c.value for c in ws[1]]
    assert headers == IDENTITY + [c["name"] for c in SHEETS["diversity"]["columns"]]
    assert any('"Employee,Worker"' in dv.formula1 for dv in ws.data_validations.dataValidation)
    assert wb.sheetnames[1] == "Instructions"
    # The blank template is recognised as the right template, just without rows
    p = parse_sheet("diversity", "t.xlsx", template_workbook("diversity"))
    assert [e["message"] for e in p.errors] == ["The file has headers but no data rows."]


def test_validation_reports_every_problem_with_row_and_column():
    rows = [_month(2024, 4, **{ENERGY: 1000}),
            {**_month(2024, 5, **{ENERGY: "lots"}), "Start_Date": "not a date"},
            _month(2024, 6, **{ENERGY: -5}),
            {**_month(2024, 7, **{ENERGY: 10}), "Location": None},
            {"Relevant_metric_energy": "Employees"}]  # filler row dragged down: skipped
    p = parse_sheet("energy", "e.xlsx", _book("energy", rows))
    found = {(e["row"], e["column"]) for e in p.errors}
    assert {(3, "Start_Date"), (3, ENERGY), (4, ENERGY), (5, "Location")} <= found
    assert any("1 row(s) had no dates" in w["message"] for w in p.warnings)
    # Header matching ignores case and punctuation; a foreign file is rejected outright
    # A file for another sheet is refused outright, not accepted with warnings
    wrong = parse_sheet("water", "e.xlsx", _book("energy", [_month(2024, 4, **{ENERGY: 5})]))
    assert not wrong.ok and "doesn't look like the Water template" in wrong.errors[0]["message"]
    assert parse_sheet("energy", "e.xlsx", _book("energy", [_month(2024, 4, **{ENERGY: 5})])).ok


def test_upload_review_extract_and_brsr(client, session, tmp_path):
    _set_storage(tmp_path)
    co, up, appr = _world(session)
    h = auth(up)

    # Catalogue and template endpoints
    cat = client.get("/api/v1/datasheets", headers=h).json()
    assert len(cat["sheets"]) == 27 and all(s["available"] for s in cat["sheets"])
    r = client.get("/api/v1/datasheets/energy/template", headers=h)
    assert r.status_code == 200 and r.headers["content-type"] == XLSX
    assert client.get("/api/v1/datasheets/nope", headers=h).status_code == 404

    good = _book("energy", [_month(2024, m, **{ENERGY: 1000 * m, "Renewable_Purchased_Heat_KWh": 100})
                            for m in range(4, 13)])
    bad = _book("energy", [_month(2024, 4, **{ENERGY: "abc"})])

    # Dry-run check: nothing stored
    r = client.post("/api/v1/datasheets/energy/check", files={"file": ("bad.xlsx", bad, XLSX)}, headers=h)
    assert r.status_code == 200 and r.json()["ok"] is False and r.json()["errors"][0]["column"] == ENERGY

    ds, v = _dataset(client, co, up, "energy")
    r = _put(client, up, ds, v, "bad.xlsx", bad)
    assert r.status_code == 422 and r.json()["error"]["code"] == "datasheet_invalid"
    assert r.json()["error"]["details"]["errors"][0] == {"row": 2, "column": ENERGY, "message": "'abc' is not a number"}
    assert _put(client, up, ds, v, "energy.xlsx", good).status_code == 201
    # The same months again in a second file would be counted twice: refused
    r = _put(client, up, ds, v, "again.xlsx", _book("energy", [_month(2024, 4, **{ENERGY: 1})]))
    assert r.status_code == 409 and r.json()["error"]["code"] == "datasheet_overlap"
    # A different month in a second file is fine
    assert _put(client, up, ds, v, "jan.xlsx", _book("energy", [_month(2025, 1, **{ENERGY: 500})])).status_code == 201

    assert client.post(f"/api/v1/datasets/{ds}/versions/{v}/submit", headers=h).status_code == 200
    from app.models.company import Company
    org_id = session.get(Company, co.id).organization_id
    from app.models.organization import Organization
    org = session.get(Organization, org_id)
    dadmin = make_user(session, "da@d.com", "deloitte", "Administrator", org=org)
    drev = make_user(session, "dr@d.com", "deloitte", "Reviewer", org=org)
    _approve(client, session, dadmin, drev, ds, v)
    _run_kpi_job(session, v)

    vals = session.exec(select(KpiValue).where(KpiValue.kpi_code == f"energy.{ENERGY}")).all()
    assert len(vals) == 10 and {kv.reporting_period_start.month for kv in vals} == {4, 5, 6, 7, 8, 9, 10, 11, 12, 1}
    kv = next(x for x in vals if x.reporting_period_start == date(2024, 6, 1))
    assert kv.value == 6000 and kv.unit == "kWh" and kv.attributes["location"] == "Mumbai"
    assert kv.reporting_period_end == date(2024, 6, 30) and kv.source_row_number == 4

    # Coverage and the BRSR figure built from it
    cov = client.get("/api/v1/datasheets/coverage?fy=2024-25", headers=h).json()
    assert cov["coverage"]["energy"]["months"] == 10
    assert any(w["sheet"] == "energy" and "10 of 12 months" in w["message"] for w in cov["warnings"])
    r = client.post("/api/v1/brsr/response-sets", json={"financial_year": "2024-25"}, headers=h)
    set_id = r.json()["public_id"]
    p6 = client.get(f"/api/v1/brsr/response-sets/{set_id}/sections/P6", headers=h).json()
    kwh = sum(1000 * m for m in range(4, 13)) + 500
    assert p6["answers"]["P6.E1"]["computed"]["non_elec"]["cy_v"] == pytest.approx(kwh * 0.0036)
    assert p6["answers"]["P6.E1"]["computed"]["ren_other"]["cy_v"] == pytest.approx(900 * 0.0036)
    assert p6["answers"]["P6.E1"]["computed"]["non_other"]["cy_v"] == 0  # reported sheet, blank column
    assert "energy" in p6["data_sheets"]["sheets"]


def test_assurance_goes_stale_when_an_approved_sheet_changes(session):
    from app.assurance.service import assured_content_hash
    from app.models.brsr import BrsrResponseSet
    from app.models.dataset import DatasetVersion
    from tests import brsr_fixture
    co, up, _ = _world(session)
    brsr_fixture.seed_sheet_data(session, co, up)
    rs = BrsrResponseSet(company_id=co.id, financial_year="2024-25", bank_version="x", created_by=up.id)
    session.add(rs)
    session.commit()
    before = assured_content_hash(session, rs)
    assert assured_content_hash(session, rs) == before  # deterministic
    kv = session.exec(select(KpiValue).where(KpiValue.kpi_code == "energy.Nonrenewable_Purchased_Electricity_from_Grid_kWh",
                                             KpiValue.reporting_period_start >= date(2024, 4, 1))).first()
    v2 = DatasetVersion(dataset_id=kv.dataset_id, version_number=2, status="approved", uploaded_by=up.id)
    session.add(v2)
    session.flush()
    session.add(KpiValue(dataset_id=kv.dataset_id, dataset_version_id=v2.id, source_file_id=kv.source_file_id,
                         source_row_number=kv.source_row_number, company_id=co.id, kpi_code=kv.kpi_code,
                         value=kv.value + 1, unit="kWh", attributes=kv.attributes,
                         reporting_period_start=kv.reporting_period_start, reporting_period_end=kv.reporting_period_end))
    session.commit()
    assert assured_content_hash(session, rs) != before


def test_formula_sheet_matches_the_code(client, session):
    from app.brsr.sheet_sources import FACTORS, SOURCES
    co, up, _ = _world(session)
    r = client.get("/api/v1/datasheets/formula-sheet", headers=auth(up))
    assert r.status_code == 200 and r.headers["content-type"] == XLSX
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Read me", "Factors", "Deviations", "Figures", "Typed in BRSR", "Questionnaire"]
    assert wb["Factors"].max_row - 1 == len(FACTORS)
    figures = {row[0] for row in wb["Figures"].iter_rows(min_row=2, values_only=True) if row[0]}
    assert set(SOURCES) <= figures
    typed = {row[0] for row in wb["Typed in BRSR"].iter_rows(min_row=2, values_only=True)}
    assert {"A.24", "P6.E7"} <= typed
